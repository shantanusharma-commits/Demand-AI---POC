"""The answer key for each golden set, from the engine run (run_golden_v2.mts), with every account's design checked
against what the engine did. Fails if any design expectation is not met.
Usage: python3 build_answer_keys.py <output folder> <spec.json> <results.json>"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT, spec, res = sys.argv[1], json.load(open(sys.argv[2])), json.load(open(sys.argv[3]))
B, WRAP, HEAD = Font(bold=True), Alignment(wrap_text=True, vertical='top'), PatternFill('solid', fgColor='DDEBF7')
OK, BAD = PatternFill('solid', fgColor='E2EFDA'), PatternFill('solid', fgColor='F8CBAD')
STEPS = {
    'A': [('Prospecting', 'Upload Golden_A_Scoring_1_Prospects.xlsx.', '17 contacts in 16 accounts. 2 issues: one opt-out (A12, never contacted) and one blank consent basis (A13, not contactable).'),
          ('Scoring', 'Set the score date to 30 Oct 2026, upload Golden_A_Scoring_2_Signals.xlsx and score.', '18 signal rows: 12 qualified; 3 rejected (stale A07, newsletter A09, company not in the list A15); 1 knock-out (A10); 2 not used (A10 excluded, A11 below fit floor).'),
          ('Scoring', 'Compare each account and person with the Expected sheet.', 'Tiers, fit and scores as listed. Score = fit x timing / 100; timing = stacked signal weights / 200.'),
          ('Hand-off', 'Hand off the default selection.', '8 prospects handed off; 1 held back (A13, no consent basis).')],
    'B': [('Prospecting', 'Upload Golden_B_Segments_NBA_1_Prospects.xlsx.', '20 contacts in 19 accounts. 1 note: B03 has no email (LinkedIn only).'),
          ('Scoring', 'Set the score date to 30 Oct 2026, upload Golden_B_Segments_NBA_2_Signals.xlsx and score.', '18 accounts in tier A or B; B19 is tier C (email click only).'),
          ('Hand-off', 'Hand off the default selection.', '18 prospects handed off, 0 held back. B19 (tier C) is not handed off.'),
          ('NBA', 'Build micro-segments.', 'Five segments form: Inquiry 4, Modernisation 3, Project 3, Engagement 3, Service renewal 3. Leadership does not form (only 2 accounts).'),
          ('NBA', 'Open each item and compare with the Expected and Draft text sheets.', 'Action, channel and claim as listed. 16 go ahead on their own; B13 and B14 go to the queue (segment too small).')],
    'C': [('Prospecting', 'Upload Golden_C_For_Review_1_Prospects.xlsx.', '17 contacts in 16 accounts. C08 is flagged: email not verified and no LinkedIn URL.'),
          ('Scoring', 'Set the score date to 30 Oct 2026, upload Golden_C_For_Review_2_Signals.xlsx and score.', 'All 16 accounts in tier A or B.'),
          ('Hand-off', 'Hand off the default selection.', '15 prospects handed off; 1 held back (C08, no way to reach them).'),
          ('NBA', 'Build micro-segments.', 'Inquiry 5, Modernisation 3, Project 3, Service renewal 3. C09 (Leadership) has no segment.'),
          ('For Review', 'Open the exception queue (each rep) and Approvals (sales manager).', 'Queue 4: C04 low confidence, C05 below deal size, C06 claim can\'t be verified, C09 segment too small. Approvals 2: C03 and C10 (sensitive).'),
          ('For Review', 'Open the released items.', '9 go ahead on their own, including C07 at version 2 (rewritten once after a pricing flag). Some may be drawn for the weekly spot-check (about 20%, at least one per rep).')],
}


def sheet(wb, title, cols, rows, widths, fills=None):
    ws = wb.create_sheet(title); ws.append(cols)
    for c in ws[1]: c.font, c.fill = B, HEAD
    for i, r in enumerate(rows):
        ws.append(r)
        if fills and fills[i]:
            for c in ws[ws.max_row]: c.fill = fills[i]
    for row in ws.iter_rows():
        for c in row: c.alignment = WRAP
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'


def check(key, acc, run):
    """Each design expectation against the engine: (what, expected, got, ok)."""
    e, out = acc['expect'], []
    a = next((x for x in run['accounts'] if x['name'] == acc['name']), None)
    rec = next((x for x in run['recs'] if x['account'] == acc['name']), None)
    if 'tier' in e: out.append(('Account tier', e['tier'], a and a['tier'], bool(a) and a['tier'] == e['tier']))
    if 'score' in e: out.append(('Account score', e['score'], a and a.get('score'), bool(a) and a.get('score') == e['score']))
    if 'scoreBelow' in e: out.append(('Score below', e['scoreBelow'], a and a.get('score'), bool(a) and (a.get('score') or 0) < e['scoreBelow']))
    if 'people' in e:
        for n, t in e['people'].items():
            p = next((p for p in a['people'] if p['name'] == n), None)
            out.append((f'{n} tier', t, p and p['tier'], bool(p) and p['tier'] == t))
    if e.get('optOut'): out.append(('Opted-out person not a prospect', 'not listed', 'not listed' if not any(p['account'] == acc['name'] and p['name'].startswith('Ana') for p in run['prospects']) else 'listed',
                                    not any(p['account'] == acc['name'] and p['name'].startswith('Ana') for p in run['prospects'])))
    if e.get('heldBack'): out.append(('Held back at hand-off', 'yes', 'yes' if any(n.endswith(acc['id']) for n in run['heldBackNames']) else 'no', any(n.endswith(acc['id']) for n in run['heldBackNames'])))
    if e.get('unmatched'): out.append(('Signal unmatched', 'Rejected', next((s['status'] for s in run['signals'] if s['company'] == 'Oakridge Chemicals'), None),
                                       any(s['company'] == 'Oakridge Chemicals' and s['status'] == 'Rejected' for s in run['signals'])))
    if 'segment' in e: out.append(('Segment', e['segment'] or 'none', rec and (rec['segment'] or 'none'), bool(rec) and rec['segment'] == e['segment']))
    if 'action' in e: out.append(('Action', e['action'], rec and rec['action'], bool(rec) and rec['action'] == e['action']))
    if 'channel' in e: out.append(('Channel', e['channel'], rec and rec['channel'], bool(rec) and rec['channel'] == e['channel']))
    if 'contact' in e: out.append(('Contact', e['contact'], rec and rec['contact'], bool(rec) and rec['contact'] == e['contact']))
    if e.get('handoff') is False: out.append(('Handed off', 'no', 'yes' if rec else 'no', rec is None))
    if 'route' in e: out.append(('Route', e['route'], rec and rec['route'], bool(rec) and rec['route'] == e['route']))
    if 'reason' in e: out.append(('Reason contains', e['reason'], rec and '; '.join(rec['reasons']), bool(rec) and any(e['reason'].lower() in t.lower() for t in rec['reasons'])))
    if e.get('regen'): out.append(('Rewritten once', 'yes', rec and (rec['regen'] or 'no'), bool(rec) and bool(rec['regen'])))
    return out


fails = 0
for key, s in spec['sets'].items():
    run = res[key]; wb = Workbook()
    sheet(wb, 'How to test', ['Step', 'Where', 'Do', 'Expect'], [[i + 1, *t] for i, t in enumerate(STEPS[key])], [6, 14, 60, 90])
    rows, crows, cfill = [], [], []
    for acc in s['accounts']:
        a = next((x for x in run['accounts'] if x['name'] == acc['name']), None)
        rec = next((x for x in run['recs'] if x['account'] == acc['name']), None)
        people = '; '.join(f"{p['name']}: {p['tier']}{'' if p.get('score') is None else ' ' + str(p['score'])}" for p in (a['people'] if a else []))
        handed = 'held back' if any(n.endswith(acc['id']) for n in run['heldBackNames']) else ('yes' if rec or any(n.endswith(acc['id']) for n in run['handoff']) else 'no')
        base = [acc['id'], acc['name'], acc['tests'], a and a.get('fit'), a and a['tier'], a and a.get('score'), people]
        if key == 'A':
            rows.append(base + [handed])
        else:
            rows.append(base + [handed, rec and (rec['segment'] or 'none'), rec and rec['action'], rec and rec['channel'], rec and ('task for the owner' if rec['task'] else rec['claim']),
                                rec and ('yes: ' + rec['regen'] if rec['regen'] else 'no'), rec and rec['route'], rec and '; '.join(rec['reasons'])])
        for what, exp, got, ok in check(key, acc, run):
            crows.append([acc['id'], what, str(exp), str(got), 'PASS' if ok else 'FAIL']); cfill.append(OK if ok else BAD); fails += not ok
    cols = ['ID', 'Account', 'What it tests', 'Fit', 'Account tier', 'Account score', 'People: tier and score', 'Handed off']
    if key == 'A':
        sheet(wb, 'Expected', cols, rows, [5, 24, 60, 6, 13, 8, 30, 11])
        sheet(wb, 'Signal rows', ['Row', 'Company', 'Signal type', 'Status', 'Weight', 'Reason'],
              [[x['row'], x['company'], x['type'], x['status'], x.get('weight'), x.get('reason') or ''] for x in run['signals']], [6, 28, 30, 18, 8, 60])
    else:
        sheet(wb, 'Expected', cols + ['Segment', 'Action', 'Channel', 'Claim (starting pack)', 'Rewritten once', 'Route', 'Reasons'], rows,
              [5, 22, 50, 6, 10, 8, 22, 9, 15, 30, 9, 14, 14, 10, 34])
        sheet(wb, 'Draft text', ['Account', 'Contact', 'Route', 'Draft, note or task'], [[r['account'], r['contact'], r['route'], r['text']] for r in run['recs']], [24, 14, 11, 110])
    sheet(wb, 'File issues', ['File', 'Row | field | issue | outcome'], [['Prospects', i] for i in run['leadIssues']] + [['Signals', i] for i in run['signalIssues']], [12, 120])
    sheet(wb, 'Engine check', ['ID', 'Check', 'Designed', 'Engine result', 'Result'], crows, [6, 30, 40, 50, 9], cfill)
    del wb['Sheet']
    wb.save(os.path.join(OUT, s['name'], f"{s['name']}_3_Answer_key.xlsx"))
    print(key, len(crows), 'checks')
print('FAILS', fails)
sys.exit(1 if fails else 0)
