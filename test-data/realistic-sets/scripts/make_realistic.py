"""Round 2 datasets: realistic client-style lists, each a workbook with a "Lead template" and a "Signal template" sheet.
Unlike the rule-by-rule sets, these look like real exports: many accounts, the usual gaps, and a business story per set.
Usage: python3 make_realistic.py <output folder> <spec.json>"""
import json, os, random, sys
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font

OUT, SPEC = sys.argv[1], sys.argv[2]
LEAD_COLS = ['first_name', 'last_name', 'job_title', 'seniority', 'email', 'email_verified', 'linkedin_url',
             'company_name', 'company_domain', 'country', 'city', 'industry', 'annual_revenue_usd', 'employee_count',
             'account_type', 'existing_customer', 'managed_separately', 'site_name', 'site_process_type', 'site_capacity',
             'capacity_unit', 'consent_basis', 'opt_out', 'owner_email']
SIG_COLS = ['company_name', 'contact_email_or_linkedin', 'signal_type', 'event_date', 'detail', 'source', 'product']
REPS = ['rep.a@client-sample.com', 'rep.b@client-sample.com']
slug = lambda v: ''.join(ch for ch in v.lower() if ch.isalnum())

CITIES = {'Singapore': ['Jurong Island', 'Tuas', 'Pulau Bukom'], 'Malaysia': ['Kerteh', 'Pengerang', 'Melaka', 'Bintulu', 'Port Dickson'],
          'Thailand': ['Map Ta Phut', 'Sriracha', 'Rayong'], 'Philippines': ['Batangas', 'Bataan', 'Limay'],
          'Indonesia': ['Cilegon', 'Balikpapan', 'Dumai', 'Bontang', 'Cilacap'], 'Vietnam': ['Dung Quat', 'Nghi Son', 'Vung Tau'],
          'Australia': ['Geelong', 'Kwinana'], 'India': ['Jamnagar', 'Kochi'], 'Japan': ['Kawasaki']}
PILOT = ['Singapore', 'Malaysia', 'Thailand', 'Philippines', 'Indonesia', 'Vietnam']
# industry, site process, capacity unit, name word
KINDS = {'refining': ('Refining', 'Crude refining', 'bpd', ['Refining', 'Refinery', 'Petroleum Refining']),
         'petchem': ('Petrochemicals', 'Olefins', 'tpa', ['Petrochem', 'Olefins', 'Aromatics', 'Polymers']),
         'lng': ('LNG', 'LNG liquefaction', 'tpa', ['LNG', 'Gas Liquefaction']),
         'gas': ('Gas processing', 'Gas processing', 'mmscfd', ['Gas Processing', 'Gas']),
         'chem': ('Specialty chemicals', 'Batch chemicals', 'tpa', ['Chemicals', 'Specialty Chemicals']),
         'power': ('Power generation', '', '', ['Power', 'Energy'])}
PREFIX = ['Amaranth', 'Banyan', 'Calamansi', 'Durian', 'Equator', 'Frangipani', 'Gaharu', 'Hornbill', 'Ironwood', 'Jasmine', 'Kenari', 'Lotus',
          'Mangrove', 'Nipah', 'Orchard', 'Pandan', 'Quartz', 'Rambutan', 'Sago', 'Teak', 'Ulin', 'Vanda', 'Waringin', 'Xylo', 'Yellowfin', 'Zircon',
          'Andaman', 'Borneo', 'Celebes', 'Mekong', 'Malacca', 'Sunda', 'Java Sea', 'Sulu', 'Natuna', 'Riau', 'Penang', 'Tioman', 'Langkawi', 'Lombok',
          'Coral', 'Monsoon', 'Typhoon', 'Tropic', 'Lagoon', 'Kingfisher', 'Seahawk', 'Tidal', 'Saltwater', 'Harbourfront', 'Delta', 'Estuary',
          'Summit', 'Highland', 'Granite', 'Basalt', 'Obsidian', 'Jade', 'Pearl', 'Sapphire', 'Topaz', 'Amber', 'Onyx', 'Opal', 'Ruby', 'Garnet',
          'Peridot', 'Beryl', 'Cobalt', 'Nickel', 'Zinc', 'Copper', 'Silver', 'Golden', 'Bronze', 'Ivory', 'Ebony', 'Cedar', 'Meranti', 'Rattan',
          'Bamboo', 'Palmyra', 'Sandalwood', 'Cinnamon', 'Clove', 'Nutmeg', 'Pepper', 'Cardamom', 'Galangal', 'Lemongrass', 'Tamarind', 'Kaffir',
          'Starfruit', 'Longan', 'Lychee', 'Mangosteen', 'Jackfruit', 'Papaya', 'Guava', 'Soursop']
SUFFIX = ['', ' Holdings', ' Industries', ' Energy', ' Group', ' Corporation', ' Bhd', ' Pte Ltd', ' Public Co', ' Tbk', ' Inc', ' JSC']
FIRST = ['Aditya', 'Afiq', 'Aiko', 'Amir', 'Anh', 'Arjun', 'Ayu', 'Bayu', 'Bich', 'Camille', 'Chandra', 'Chen', 'Darius', 'Dewi', 'Duc', 'Elaine',
         'Erwin', 'Farid', 'Gita', 'Hafiz', 'Hana', 'Hiroshi', 'Indra', 'Iris', 'Jasmine', 'Joaquin', 'Kartika', 'Khanh', 'Kiran', 'Lakshmi', 'Lan',
         'Lukas', 'Mai', 'Malik', 'Marisol', 'Mei', 'Minh', 'Nadia', 'Naveen', 'Nur', 'Pablo', 'Pham', 'Putri', 'Qistina', 'Rafael', 'Rahul',
         'Rizal', 'Rohan', 'Sakura', 'Sani', 'Siti', 'Somsak', 'Sunil', 'Tariq', 'Thanh', 'Uday', 'Vanessa', 'Wei', 'Wira', 'Xin', 'Yasmin',
         'Yuto', 'Zainab', 'Zhi']
LAST = ['Abdullah', 'Aquino', 'Bakar', 'Bui', 'Castillo', 'Chandra', 'Cheong', 'Dang', 'Dela Cruz', 'Ferrer', 'Gonzaga', 'Hadi', 'Hamid',
        'Hartono', 'Ho', 'Ishak', 'Iyer', 'Kamal', 'Kurniawan', 'Lai', 'Lim', 'Mahmud', 'Manalo', 'Menon', 'Nair', 'Ng', 'Nguyen', 'Ong', 'Pang',
        'Pillai', 'Pratama', 'Quah', 'Rahman', 'Reyes', 'Santoso', 'Setiawan', 'Sharma', 'Sim', 'Siregar', 'Suryadi', 'Tan', 'Teo', 'Thapa',
        'Tran', 'Valdez', 'Vo', 'Wahyudi', 'Wong', 'Yap', 'Yusof', 'Zulkarnain']
TITLES = {'primary': [('Instrumentation Manager', 'Manager'), ('Head of Automation', 'Director'), ('Control Systems Engineer', 'Senior'),
                      ('I&C Lead', 'Manager'), ('DCS Engineer', 'Senior'), ('E&I Superintendent', 'Manager'), ('Automation Engineer', 'Senior')],
          'secondary': [('Plant Manager', 'Director'), ('Operations Director', 'Director'), ('Maintenance Manager', 'Manager'),
                        ('Reliability Engineer', 'Senior'), ('Procurement Manager', 'Manager'), ('Project Director', 'Director')],
          'other': [('Finance Manager', 'Manager'), ('HR Business Partner', 'Senior'), ('Sustainability Lead', 'Manager'),
                    ('Commercial Manager', 'Manager'), ('IT Manager', 'Manager'), ('Legal Counsel', 'Senior')]}
SIGNALS = [  # type, detail, source, oldest age it is drawn at
    ('Inquiry or RFQ', ['Sent an RFQ for a DCS migration', 'Asked for budget pricing on a safety system upgrade', 'Requested a call about alarm management',
                        'Asked about a control room modernisation'], 'Web form', 100),
    ('Installed system near end of support', ['DCS on the crude unit reaches end of support in 2027', 'Legacy PLCs at the site are end of life',
                                              'Safety system goes out of support next year'], 'Installed base', 800),
    ('Service contract renewal', ['Service contract renews on 31 Mar 2027', 'Annual maintenance agreement renews in Q1'], 'CRM', 200),
    ('Capital project', ['Announced a site expansion', 'FEED awarded for a new unit', 'Final investment decision on a new train'], 'News', 400),
    ('Leadership change', ['New head of operations appointed', 'New plant director appointed', 'New VP engineering joined'], 'News', 140),
    ('Webinar attended', ['Attended the legacy-control migration webinar', 'Attended the alarm management webinar'], 'Webinar platform', 200),
    ('Webinar registered', ['Registered for the November reliability webinar', 'Registered for the cyber-resilience webinar'], 'Webinar platform', 130),
    ('Content download', ['Downloaded the migration planning guide', 'Downloaded the alarm rationalisation checklist'], 'Website', 200),
    ('Email clicked', ['Clicked the lifecycle services email', 'Clicked the modernisation case-study email'], 'Email platform', 70),
]
SIG_BY = {s[0]: s for s in SIGNALS}

def build(spec):
    """spec keys: id, title, story, n, seed, asOf, kinds (weights), existing, pilot share, contacts (min,max), signal mix, mess (probabilities)."""
    rnd = random.Random(spec['seed'])
    as_of = date.fromisoformat(spec['asOf'])
    m = spec.get('mess', {})
    p = lambda k: rnd.random() < m.get(k, 0)
    leads, sigs, notes, used_names = [], [], [], set()
    for i in range(spec['n']):
        kind = rnd.choices(list(spec['kinds']), weights=list(spec['kinds'].values()))[0]
        ind, proc, unit, words = KINDS[kind]
        while True:
            name = f"{rnd.choice(PREFIX)} {rnd.choice(words)}{rnd.choice(SUFFIX)}"
            if name not in used_names: used_names.add(name); break
        country = rnd.choice(PILOT) if rnd.random() < spec.get('pilot', 0.95) else rnd.choice(['Australia', 'India', 'Japan'])
        dom = slug(name.split(' ')[0] + name.split(' ')[1]) + '.example'
        rev = rnd.choice(spec.get('revenue', [4.5e9, 2.2e9, 1.1e9, 6e8, 3e8, 1.5e8]))
        emp = int(rev / 6e5) if rnd.random() < 0.8 else ''
        if p('no_size'): rev, emp = '', ''
        elif p('revenue_text'): rev = f"{int(rev):,}"
        existing = 'Y' if rnd.random() < spec.get('existing', 0.3) else 'N'
        acct = dict(company_name=name, company_domain=dom, country=country, city=rnd.choice(CITIES[country]), industry=ind,
                    annual_revenue_usd=int(rev) if isinstance(rev, float) else rev, employee_count=emp,
                    account_type=rnd.choices(['Owner-operator', 'EPC contractor', 'System integrator', 'Distributor'], weights=spec.get('types', [78, 10, 6, 6]))[0],
                    existing_customer=existing, managed_separately='Y' if p('managed') else 'N',
                    site_name=f"{name.split(' ')[0]} {'site' if not proc else proc.split()[0].lower() + ' plant'}" if proc else '',
                    site_process_type=proc, site_capacity=rnd.choice([80000, 150000, 250000, 400000]) if proc else '', capacity_unit=unit,
                    owner_email='' if p('no_owner') else REPS[i % 2])
        if p('unmapped'): acct['industry'] = rnd.choice(['Water utilities', 'Cement', 'Mining', 'Food processing'])
        k = rnd.randint(*spec.get('contacts', (2, 5)))
        roles = ['primary' if rnd.random() < spec.get('primary', 0.85) else 'secondary'] + rnd.choices(['primary', 'secondary', 'other'], weights=[30, 45, 25], k=k - 1)
        people, seen = [], set()
        for role in roles:
            while True:
                fn, ln = rnd.choice(FIRST), rnd.choice(LAST)
                if (fn, ln) not in seen: seen.add((fn, ln)); break
            title, sen = rnd.choice(TITLES[role])
            row = dict(acct, first_name=fn, last_name=ln, job_title=title, seniority=sen,
                       email=f"{slug(fn)}.{slug(ln)}@{dom}", email_verified='Y',
                       linkedin_url=f"https://www.linkedin.com/in/{slug(fn)}-{slug(ln)}-{rnd.randint(100, 999)}",
                       consent_basis='existing customer' if existing == 'Y' else rnd.choice(['legitimate interest', 'legitimate interest', 'opted in']), opt_out='N')
            if p('unverified'): row['email_verified'] = 'N'
            if p('no_linkedin'): row['linkedin_url'] = ''
            if p('opt_out'): row['opt_out'] = 'Y'
            if p('no_consent'): row['consent_basis'] = ''
            if p('bad_email'): row['email'] = row['email'].replace('.example', '')
            if p('yes_no_words'): row['email_verified'] = {'Y': 'Yes', 'N': 'No'}[row['email_verified']]
            leads.append(row); people.append(row)
        if p('duplicate_row'):
            d = dict(people[0]); d['email'] = d['email'].upper(); leads.append(d)
        # Signals: how many, which types, how old
        n_sig = rnd.choices([0, 1, 2, 3, 4], weights=spec.get('signal_count', [30, 35, 20, 10, 5]))[0]
        for _ in range(n_sig):
            typ = rnd.choices([s[0] for s in SIGNALS], weights=spec['mix'])[0]
            _, details, src, oldest = SIG_BY[typ]
            if existing == 'N' and typ in ('Service contract renewal', 'Installed system near end of support') and rnd.random() < 0.7:
                typ = rnd.choice(['Webinar attended', 'Content download', 'Capital project']); _, details, src, oldest = SIG_BY[typ]
            who = rnd.choice(people)
            ref = who['email'] if who['email'] and rnd.random() < 0.85 else (who['linkedin_url'] or who['email'])
            age = rnd.randint(0, int(oldest * spec.get('age_share', 1)))
            if spec.get('safe_details'): details = [x for x in details if 'cyber' not in x and 'pricing' not in x]
            when = as_of - timedelta(days=age)
            sigs.append([name, '' if p('no_person') else ref, typ, when, rnd.choice(details), src, ''])
        if p('knockout'):
            sigs.append([name, '', 'Installed system (current)', as_of - timedelta(days=rnd.randint(100, 900)), 'Current pilot product installed on the main unit', 'Installed base', 'Current platform'])
    # Signal-file noise: rows a real export would carry
    extra = []
    for s in sigs:
        if p('sig_unknown_person') and s[1]: extra.append([s[0], 'someone.new@' + s[1].split('@')[-1] if '@' in s[1] else s[1] + 'x', s[2], s[3], s[4], s[5], ''])
        if p('sig_dup'): extra.append(list(s))
        if p('sig_newsletter'): extra.append([s[0], s[1], 'Newsletter sign-up', s[3], 'Signed up for the newsletter', 'Website', ''])
    sigs += extra
    for _ in range(int(m.get('sig_unknown_company', 0) * len(sigs))):
        sigs.append([f"{rnd.choice(PREFIX)} Unlisted Refining", '', 'Inquiry or RFQ', as_of - timedelta(days=rnd.randint(1, 60)), 'Asked for a quote', 'Web form', ''])
    rnd.shuffle(sigs)
    return leads, sigs

def write(spec, leads, sigs, path, prior=None):
    wb = Workbook()
    ws = wb.active; ws.title = 'Lead template'
    hdr = spec.get('lead_headers') or LEAD_COLS
    pre = spec.get('title_row')
    if pre: ws.append([pre])
    ws.append(hdr + spec.get('extra_lead_cols', []))
    for i, r in enumerate(leads):
        extra = [f'LD-{100000 + i}', '2026-0' + str(1 + i % 9) + '-15', ['Trade show', 'Website', 'Referral', 'Purchased list'][i % 4]][:len(spec.get('extra_lead_cols', []))]
        ws.append([None if r.get(c, '') == '' else r.get(c) for c in LEAD_COLS] + extra)
    ws2 = wb.create_sheet('Signal template')
    ws2.append((spec.get('sig_headers') or SIG_COLS) + spec.get('extra_sig_cols', []))
    for i, s in enumerate(sigs):
        d = s[3]
        val = datetime(d.year, d.month, d.day) if spec.get('excel_dates') else d.isoformat()
        ws2.append([None if v == '' else v for v in s[:3]] + [val] + [None if v == '' else v for v in s[4:]] + [f'ACT-{500000 + i}'][:len(spec.get('extra_sig_cols', []))])
        if spec.get('excel_dates'): ws2.cell(ws2.max_row, 4).number_format = 'dd-mmm-yyyy'
    for w in (ws, ws2):
        hr = 2 if (w is ws and pre) else 1
        for c in w[hr]: c.font = Font(bold=True)
        for col in w.columns: w.column_dimensions[col[0].column_letter].width = min(38, max(10, *(len(str(c.value or '')) + 2 for c in col[:200])))
        w.freeze_panes = f'A{hr + 1}'
    wb.save(path)

BASE_MIX = [8, 10, 7, 9, 7, 16, 10, 16, 17]
SETS = [
    dict(id='R1', file='R1_Smoke_test_small_clean.xlsx', title='Smoke test: small and clean', n=12, seed=101, asOf='2026-10-20', kinds={'refining': 8, 'petchem': 2, 'lng': 1, 'gas': 1},
         existing=0.5, pilot=1.0, contacts=(2, 4), primary=1.0, revenue=[4.5e9, 2.2e9, 1.4e9], types=[100, 0, 0, 0], age_share=0.3, safe_details=True,
         mix=[30, 25, 0, 20, 0, 25, 0, 0, 0], signal_count=[0, 50, 40, 10, 0],
         story='A quick daily check after every fix: 12 accounts, nothing wrong in the data. Run it end to end in 10 minutes; everything should go through without a flag.'),
    dict(id='R2', file='R2_Pilot_list_week_1.xlsx', title='Pilot list, week 1', n=80, seed=202, asOf='2026-10-20', kinds={'refining': 34, 'petchem': 24, 'lng': 10, 'gas': 12, 'chem': 12, 'power': 8},
         existing=0.3, pilot=0.94, mix=BASE_MIX,
         mess=dict(unverified=0.06, no_linkedin=0.12, opt_out=0.04, no_consent=0.03, no_size=0.05, unmapped=0.05, no_owner=0.03, managed=0.03, duplicate_row=0.05,
                   bad_email=0.02, knockout=0.04, sig_unknown_person=0.03, sig_dup=0.03, sig_newsletter=0.03, sig_unknown_company=0.03, no_person=0.05),
         story='What the first real pilot file is likely to look like: 80 accounts across the region and all verticals, with the usual 5-10% of gaps. Use it for the main end-to-end run and the For Review walk-through.'),
    dict(id='R3', file='R3_Large_list_performance.xlsx', title='Large list for performance', n=500, seed=303, asOf='2026-10-20', kinds={'refining': 30, 'petchem': 25, 'lng': 10, 'gas': 12, 'chem': 15, 'power': 8},
         existing=0.25, pilot=0.92, mix=BASE_MIX,
         mess=dict(unverified=0.06, no_linkedin=0.12, opt_out=0.04, no_consent=0.03, no_size=0.05, unmapped=0.05, no_owner=0.03, managed=0.02, duplicate_row=0.04,
                   knockout=0.03, sig_unknown_person=0.02, sig_dup=0.02, sig_newsletter=0.02, sig_unknown_company=0.02, no_person=0.05),
         story='Several times the pilot size: 500 accounts. Time each step (upload, scoring, building Micro-segments, opening For Review) and watch for the page freezing.'),
    dict(id='R5', file='R5_Installed_base_existing_customers.xlsx', title='Installed base: existing customers', n=30, seed=505, asOf='2026-10-20', kinds={'refining': 14, 'petchem': 8, 'lng': 4, 'gas': 4},
         existing=1.0, pilot=1.0, mix=[6, 34, 30, 6, 6, 6, 4, 4, 4], signal_count=[10, 40, 30, 15, 5], mess=dict(knockout=0.12, no_linkedin=0.1),
         story='A list of existing customers where the signals come from the installed base: end-of-support systems and service renewals. Expect Modernisation and Service renewal segments, and several accounts knocked out because the pilot product is already installed.'),
    dict(id='R6', file='R6_Net_new_marketing_engagement.xlsx', title='Net-new list from marketing', n=40, seed=606, asOf='2026-10-20', kinds={'refining': 14, 'petchem': 10, 'lng': 5, 'gas': 5, 'chem': 6},
         existing=0.0, pilot=1.0, primary=0.6, mix=[3, 0, 0, 3, 4, 25, 25, 25, 15], signal_count=[15, 40, 30, 10, 5],
         mess=dict(unverified=0.15, no_linkedin=0.2, no_consent=0.05),
         story='A list of companies Client has never sold to, warmed up by marketing: webinars, downloads and email clicks, which are weak signals. Expect mostly tier B and C, one large Engagement segment, and more low-confidence items (fewer buyers in the primary role).'),
    dict(id='R7', file='R7_CRM_export_format.xlsx', title='CRM export format', n=25, seed=707, asOf='2026-10-20', kinds={'refining': 10, 'petchem': 8, 'gas': 4, 'lng': 3},
         existing=0.3, pilot=1.0, mix=BASE_MIX, excel_dates=True, title_row='Leads export - generated by CRM',
         lead_headers=[c.replace('_', ' ').title() for c in LEAD_COLS], extra_lead_cols=['Lead ID', 'Created Date', 'Lead Source'], extra_sig_cols=['Activity ID'],
         mess=dict(yes_no_words=0.5, revenue_text=0.3, no_linkedin=0.1),
         story='The same kind of list as it comes out of a CRM rather than the template: a title row, headers like "First Name" and "Company Name", extra CRM columns, "Yes"/"No" instead of Y/N, revenue with commas, and real Excel dates in the signal sheet. Everything should still be read.'),
]

spec_out = []
os.makedirs(OUT, exist_ok=True)
for s in SETS:
    leads, sigs = build(s)
    write(s, leads, sigs, os.path.join(OUT, s['file']))
    spec_out.append(dict(id=s['id'], file=s['file'], title=s['title'], story=s['story'], asOf=s['asOf'], accounts=s['n'], leads=len(leads), signals=len(sigs)))
    print(s['id'], s['n'], 'accounts', len(leads), 'lead rows', len(sigs), 'signal rows')
    if s['id'] == 'R2':
        # R4: the same list one week later: last week's signals plus the new week's, scored a week later
        rnd = random.Random(404)
        names = sorted({r['company_name'] for r in leads})
        by_co = {}
        for r in leads: by_co.setdefault(r['company_name'], []).append(r)
        new = []
        for nm in rnd.sample(names, 22):
            typ = rnd.choices([x[0] for x in SIGNALS], weights=[16, 6, 4, 8, 6, 16, 12, 16, 16])[0]
            who = rnd.choice(by_co[nm])
            new.append([nm, who['email'] or who['linkedin_url'], typ, date(2026, 10, 27) - timedelta(days=rnd.randint(0, 6)), rnd.choice(SIG_BY[typ][1]), SIG_BY[typ][2], ''])
        s4 = dict(s, id='R4', file='R4_Pilot_list_week_2.xlsx', asOf='2026-10-27')
        write(s4, leads, sigs + new, os.path.join(OUT, s4['file']))
        spec_out.append(dict(id='R4', file=s4['file'], title='Pilot list, week 2', asOf='2026-10-27', accounts=s['n'], leads=len(leads), signals=len(sigs) + len(new),
                             story='The R2 list a week later: the same leads, last week\'s signals plus 22 new ones from 21-27 Oct, scored as of 27 Oct. Use it after working through R2 in For Review: scores move, some accounts change tier, and accounts you approved, sent or rejected as already engaged in week 1 are set aside (engage-once).',
                             newSignals=len(new), prior='R2'))
        print('R4', len(new), 'new signals')
json.dump(spec_out, open(SPEC, 'w'), indent=1)
