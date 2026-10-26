"""Adds the answer-key and how-to sheets to each Round 2 workbook, and writes Round2_Guide.xlsx.
Usage: python3 add_answer_keys.py <sets folder> <spec.json> <results.json>"""
import json, os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

d, spec, res = sys.argv[1], json.load(open(sys.argv[2])), json.load(open(sys.argv[3]))
F = 'Arial'
thin = Side(style='thin', color='FFBFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F1F1F')
CAT = {'sensitive': 'Sensitive content', 'claim': 'Unsupported claim', 'low': 'Low confidence', 'small': 'Segment too small',
       'deal': 'Below deal size', 'brand': 'Brand or rule check', 'reach': "Can't reach them", 'other': 'Other'}
TIERS = ['A', 'B', 'C', 'Watch list', 'Below fit floor', 'Excluded']

def table(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=F, size=10, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(name=F, size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths): ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions

def steps(s, r):
    st = r['stats']; i = st['issuesBySeverity']
    tiers = ', '.join(f"{t} {r['tiers'][t]}" for t in TIERS if t in r['tiers'])
    sig = ', '.join(f'{v} {k}' for k, v in sorted(r['signalStatus'].items(), key=lambda x: -x[1]))
    reasons = ', '.join(f"{CAT[k]} {v}" for k, v in sorted(r['reasons'].items(), key=lambda x: -x[1])) or 'none'
    return [
        ['1. Prospecting', f"Upload {s['file']} (reads the \"Lead template\" sheet).",
         f"{st['rowsRead']} rows read; {st['notLoaded']} not loaded; {st['merged']} merged as duplicates; {st['accounts']} accounts, {st['accountsPass']} pass; "
         f"{st['inFinalList']} contacts in the final list. Issues: {i['flag']} flagged, {i['info']} notes."],
        ['2. Account Scoring', f"Save the list (or upload the same file as the prospect file in Scoring). Score as of {s['asOf']}; drop the same file as the signal file (reads \"Signal template\").",
         f"{r['signalRows']} signal rows: {sig}. {r['suggested']} suggested new contact(s). Tiers: {tiers}."],
        ['3. Micro-segments & NBA', f"Build from this scoring, scoring date {s['asOf']}.",
         f"Segments: {', '.join(x['name'] + ' ' + str(x['n']) for x in r['segments'])}; {r['exceptions']} exception(s)."],
        ['4. For Review', 'Open For Review (and Today).', f"{r['review']} item(s) need a person, {r['ahead']} go ahead on their own. Reasons: {reasons}."],
    ]

def account_rows(r):
    out = []
    for a in r['accounts']:
        review = ('; '.join(CAT[c] for c in a['cats']) + (' (due ' + ('same day' if a['due'] == 0 else f"{a['due']} working days") + ')') if a['cats'] else 'Goes ahead') if a['rec'] else ''
        out.append([a['name'], a['country'], a['vertical'], a['owner'] or '(none)', 'Y' if a['existing'] else 'N', a['fit'] if a['fit'] is not None else '',
                    a['score'] if a['score'] is not None else '', a['tier'], a['signals'], a['confidence'] or a['why'],
                    a['segment'] or ('No segment: exception' if a['rec'] else ''), a['action'], a['contact'], a['channel'], review])
    return out
ACC_HDR = ['Account', 'Country', 'Vertical', 'Owner', 'Existing customer', 'Fit', 'Score', 'Tier', 'Qualified signals', 'Confidence, or why not scored',
           'Micro-segment', 'Recommended action', 'Contact', 'Channel', 'For Review']
ACC_W = [34, 12, 16, 26, 9, 7, 7, 13, 9, 34, 22, 36, 18, 10, 34]

guide_rows = []
for s in spec:
    r = res[s['id']]
    path = os.path.join(d, s['file'])
    wb = load_workbook(path)
    for n in ('How to use', 'Answer key - summary', 'Answer key - accounts', 'Week 1 vs week 2'):
        if n in wb.sheetnames: del wb[n]
    ws = wb.create_sheet('Answer key - summary')
    table(ws, ['Step', 'What to do', 'What you should see'], steps(s, r), [22, 60, 90])
    table(wb.create_sheet('Answer key - accounts'), ACC_HDR, account_rows(r), ACC_W)
    if s.get('prior'):
        p = {a['name']: a for a in res[s['prior']]['accounts']}
        rows = []
        for a in r['accounts']:
            b = p[a['name']]
            moved = 'Up' if TIERS.index(a['tier']) < TIERS.index(b['tier']) else 'Down' if TIERS.index(a['tier']) > TIERS.index(b['tier']) else ''
            if moved or a['score'] != b['score'] or a['segment'] != b['segment']:
                rows.append([a['name'], b['score'] if b['score'] is not None else '', b['tier'], b['segment'], a['score'] if a['score'] is not None else '',
                             a['tier'], a['segment'], moved or 'Same tier'])
        table(wb.create_sheet('Week 1 vs week 2'), ['Account', 'Week 1 score', 'Week 1 tier', 'Week 1 segment', 'Week 2 score', 'Week 2 tier', 'Week 2 segment', 'Tier moved'],
              rows, [34, 10, 12, 18, 10, 12, 18, 10])
    hw = wb.create_sheet('How to use')
    hw.column_dimensions['A'].width = 120
    lines = [f"{s['id']} {s['title']}: {s['accounts']} accounts, {s['leads']} lead rows, {s['signals']} signal rows", s['story'], '',
             'Drop this same file twice: in Prospecting (or Account Scoring > "Upload a prospect file") it reads "Lead template"; as the signal file it reads "Signal template".',
             f"Set Score as of to {s['asOf']}. Then follow \"Answer key - summary\" step by step and check single accounts in \"Answer key - accounts\".",
             'The answer keys come from reading this file the way the POC does and running it through the POC\'s own rules (legacy/demandai-engine.js).',
             'Everything in this file is invented: companies, people and domains.']
    if s['id'] == 'R4': lines.insert(3, 'Before this set: run R2 and decide its items in For Review (approve, send, or reject as "already engaged"). Those accounts must be set aside here as "Engaged once". The counts in the answer key assume nothing was decided in R2: each account you decided moves out of its segment into "Engaged once". "Week 1 vs week 2" shows how scores move.')
    if s['id'] == 'R3': lines.insert(3, 'Time each step with a stopwatch and note it in Round2_Guide.xlsx: upload, scoring, building Micro-segments, opening For Review. Note any freeze or error.')
    for i, l in enumerate(lines, 1):
        c = hw.cell(i, 1, l); c.font = Font(name=F, size=11 if i == 1 else 10, bold=i == 1); c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)
    st = r['stats']
    guide_rows.append([s['id'], s['file'], s['title'], s['story'], s['asOf'], st['accounts'], st['rowsRead'], r['signalRows'],
                       ', '.join(f"{t} {r['tiers'][t]}" for t in TIERS if t in r['tiers']), ', '.join(x['name'] + ' ' + str(x['n']) for x in r['segments']),
                       f"{r['review']} / {r['ahead']}", 'Not run', '', '', '', ''])

wb = Workbook(); ws = wb.active; ws.title = 'Datasets'
table(ws, ['Set', 'File', 'Name', 'What it is for', 'Score as of', 'Accounts', 'Lead rows', 'Signal rows', 'Expected tiers', 'Expected segments',
           'For Review / go ahead', 'Status', 'Upload time (s)', 'Scoring time (s)', 'Gap IDs', 'Notes'], guide_rows,
      [5, 34, 26, 60, 11, 9, 9, 9, 34, 34, 11, 11, 10, 10, 10, 30])
dv = DataValidation(type='list', formula1='"Not run,Pass,Fail,Blocked"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'L2:L{ws.max_row}')
for r in range(2, ws.max_row + 1):
    for c in 'LMNOP': ws[f'{c}{r}'].font = Font(name=F, size=10, color='FF0000FF')
hw = wb.create_sheet('How to use', 0); hw.column_dimensions['A'].width = 120
for i, l in enumerate([
        'Round 2 datasets: realistic client-style lists',
        'The first round (upload-test-pack, poc-test-sets) checks one rule per row. These sets look like real files instead: many accounts, the usual gaps, and a business story each.',
        'Each workbook holds both the leads ("Lead template") and the signals ("Signal template"): drop the same file in Prospecting and again as the signal file in Account Scoring.',
        '',
        'Suggested order: R1 after every fix (10 minutes) > R7 (does a CRM export read?) > R2 then R4 (the pilot weeks) > R5 and R6 (two kinds of client list) > R3 (performance).',
        'Use the Status, time and Gap ID columns in "Datasets" during the test window. Each workbook has its own answer keys and a "How to use" sheet.',
        'All companies, people and domains are invented.'], 1):
    c = hw.cell(i, 1, l); c.font = Font(name=F, size=11 if i == 1 else 10, bold=i == 1); c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(os.path.join(d, 'Round2_Guide.xlsx'))
print('done')
