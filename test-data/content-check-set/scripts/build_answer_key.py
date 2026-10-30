"""Content_Check_3_Answer_key.xlsx from the engine runs (with the Content library edits, and the starting pack).
Usage: python3 build_answer_key.py <out.xlsx> <spec.json> <with_edits.json> <starting_pack.json>"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT, spec, new, base = sys.argv[1], *(json.load(open(p)) for p in sys.argv[2:5])
NEWS = 'We have published a newsletter that help to get us 50% more revenue.'
B, WRAP, HEAD = Font(bold=True), Alignment(wrap_text=True, vertical='top'), PatternFill('solid', fgColor='DDEBF7')
CLAIM = {r['claim']: r['claimText'] for r in base['recs'] + new['recs'] if r['claim']}
wb = Workbook()

def sheet(title, cols, rows, widths):
    ws = wb.create_sheet(title); ws.append(cols)
    for c in ws[1]: c.font, c.fill = B, HEAD
    for r in rows: ws.append(r)
    for row in ws.iter_rows():
        for c in row: c.alignment = WRAP
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    return ws

steps = [
    ('1', 'Content library', 'Keep your edits live: the PPT guideline, the newsletter claim, "Just checking in" in openers to avoid, email length 150. Publish the draft if it is not live yet.', 'History shows your version as live.'),
    ('2', 'Prospecting', 'Upload Content_Check_1_Prospects.xlsx.', '20 accounts, 38 contacts, no errors. 2 contacts have a LinkedIn URL and no email.'),
    ('3', 'Scoring', 'Upload Content_Check_2_Signals.xlsx and score.', '20 signals accepted. Tier A: Inquiry, Modernisation, Project (scores 50, 50, 40). Tier B: Service renewal, Engagement (30) and Leadership (22.5).'),
    ('4', 'Hand off', 'Hand off the default selection (tier A and B with a live signal).', '20 prospects handed off, 0 held back (one per account: the person named on the signal).'),
    ('5', 'NBA', 'Build micro-segments.', 'All six segments form: Inquiry 4, Engagement 4, Modernisation 3, Project 3, Service renewal 3, Leadership 3.'),
    ('6', 'NBA / drafts', 'Open each draft and compare it with the "Expected" sheet.', 'The claim line in each email matches the "Claim with your edits" column.'),
    ('7', 'For Review', 'Open the queue and Approvals.', 'Both empty for this build: all 20 go ahead on their own. About 20% may later show up as weekly spot-checks (at least one per rep).'),
    ('8', 'Optional', 'Remove the newsletter claim (or set it to expired), publish, and build again from the same scoring run.', 'Inquiry emails go back to CS-01; the 3 Leadership items become tasks for the account owner (no claim to cite). See "Without your edits".'),
]
sheet('How to test', ['Step', 'Where', 'Do', 'Expect'], steps, [6, 16, 70, 80])

def find(d, acc): return next(r for r in d['recs'] if r['account'] == acc)
rows = []
for a in spec['accounts']:
    n, b = find(new, a['name']), find(base, a['name'])
    li = n['channel'] == 'LinkedIn'
    if li: n, b = dict(n, claim=''), dict(b, claim='')
    changed = 'Yes' if (n['claim'], n['channel']) != (b['claim'], b['channel']) else 'No'
    look = ('Newsletter claim is the proof line' if n['claim'] == 'CL-07' and n['channel'] == 'Email'
            else 'LinkedIn note: no claim is written into a note' if n['channel'] == 'LinkedIn'
            else 'Same claim as before your edits')
    rows.append([a['id'], a['name'], a['contact'], a['owner'], n['tier'], n['score'], n['segment'], n['action'], n['channel'],
                 n['claim'] or '-', CLAIM.get(n['claim'], '') if not li else 'none: a note carries no claim', b['claim'] or ('none: a note' if li else 'none: task for the owner'), changed, n['words'] or '', n['route'], look])
sheet('Expected', ['ID', 'Account', 'Contact', 'Owner', 'Tier', 'Score', 'Segment', 'Action', 'Channel', 'Claim with your edits', 'Claim text',
                   'Claim without your edits', 'Changed by your edits', 'Words', 'Route', 'What to look for'], rows,
      [5, 22, 16, 24, 6, 7, 15, 34, 9, 12, 50, 18, 10, 7, 10, 40])

sheet('Draft text', ['ID', 'Account', 'Channel', 'Subject', 'Body or note (with your edits)'],
      [[a['id'], a['name'], find(new, a['name'])['channel'], find(new, a['name'])['subject'], find(new, a['name'])['body']] for a in spec['accounts']],
      [5, 22, 9, 22, 100])
sheet('Without your edits', ['ID', 'Account', 'Channel', 'Claim', 'Body, note or task (starting pack)'],
      [[a['id'], a['name'], find(base, a['name'])['channel'], find(base, a['name'])['claim'] or '-', find(base, a['name'])['body']] for a in spec['accounts']],
      [5, 22, 9, 8, 100])

checks = [
    ('Newsletter claim', 'Added, approved at once, and its words name no topic, so it fits all six segments.',
     'Used as the proof line in the 3 Inquiry emails and 3 Leadership emails (6 emails). Other segments keep their own case study, which ranks first there.',
     'Nothing flags it: there is no claim approval step and no quality check. "help to get us" and the unsupported 50% go out as written.'),
    ('If your claim title names a topic', 'Where a claim fits is read from its text and title together (renewal, project, webinar, upgrade, question and so on).',
     'This key assumes the title names no topic, so the claim fits all six segments. If your title does name one, the claim only fits that segment.', 'Check the claim row in the Content library: it lists its segments.'),
    ('Leadership', 'Without your claim, the only Leadership claim (PS-06) expired on 2026-06-30.',
     'With it, the 3 Leadership items become real emails instead of tasks for the account owner.', 'The visible change in this build.'),
    ('"Just checking in"', 'Already in the starting list of openers to avoid.', 'No change: no rule draft opens with it.', ''),
    ('Email length 150', 'Only the length check reads it.', 'No change: rule drafts are 42 to 54 words, far under both 120 and 150. An AI draft is still briefed "under 80 words".', ''),
    ('PPT guideline', 'Stored for reference only; rules are read from .txt and .md files.', 'No suggested rules, nothing in the drafts changes.', ''),
    ('Routing', 'All 20 pass every check.', 'Queue 0, Approvals 0, Released 20.', 'Signal details carry no sensitive or pricing words on purpose.'),
    ('LinkedIn notes', 'Seagrove (Inquiry) and Tidewell (Engagement) have no email.', 'A connection note goes out; a note carries no claim, so the newsletter does not appear there.', ''),
]
sheet('What changes', ['Edit', 'How the POC treats it', 'What you see in this build', 'Note'], checks, [20, 50, 60, 55])
del wb['Sheet']; wb.save(OUT)
