"""Writes the scenario-set workbooks (TS6-TS12) and POC_Test_Cases.xlsx, the module-wise list of test cases for every set.
Usage: python3 build_scenarios.py <scenarios.json> <scenario_results.json> <ts_inputs.json> <ts_results.json> <output folder>"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

sc, scr, ts, tsr = (json.load(open(p)) for p in sys.argv[1:5])
outdir = sys.argv[5]
F = 'Arial'
thin = Side(style='thin', color='FFBFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F1F1F')
AS_OF = ts['asOf']
CAT = {'sensitive': 'Sensitive content', 'claim': 'Unsupported claim', 'low': 'Low confidence', 'small': 'Segment too small',
       'deal': 'Below deal size', 'brand': 'Brand or rule check', 'reach': "Can't reach them", 'other': 'Other'}

def sheet(wb, title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=F, size=10, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'A2'
    return ws

def guide(wb, lines):
    ws = wb.create_sheet('How to use')
    ws.column_dimensions['A'].width = 120
    for i, l in enumerate(lines, 1):
        c = ws.cell(i, 1, l); c.font = Font(name=F, size=11 if i == 1 else 10, bold=(i == 1)); c.alignment = Alignment(wrap_text=True, vertical='top')

lead_w = [12, 12, 30, 10, 34, 8, 40, 26, 26, 12, 14, 18, 14, 10, 15, 9, 9, 20, 18, 10, 8, 18, 8, 30]
sig_w = [26, 36, 34, 12, 50, 16, 14]
num = lambda v: '' if v is None else v

def review_text(a):
    if not a['rec']: return 'No item (no recommendation)'
    if not a['cats']: return 'None: goes ahead on its own (may be drawn for a spot-check)'
    due = 'same day' if a['dueDays'] == 0 else f"{a['dueDays']} working day{'s' if a['dueDays'] != 1 else ''}"
    return '; '.join(CAT[c] for c in a['cats']) + f' (due: {due})'

def after_text(a):
    if a['excluded']: return 'Excluded: ' + a['excludedReason']
    if a['tier'] == 'Below fit floor': return 'Below fit floor: not scored, not passed to Micro-segments'
    if a['tier'] == 'Watch list': return 'No qualified signal: watch list'
    if a['segment']: return a['segment'] + (' (joined on its next signal)' if a['fallback'] else '')
    if a['rec']: return 'No segment: exception in For Review'
    return 'No recommendation'

def expected_line(a):
    """One line of what to expect for an account, for the test cases."""
    parts = []
    if a['excluded']: return 'Excluded: ' + a['excludedReason']
    if a['fit'] is not None: parts.append(f"Fit {a['fit']:g}")
    if a['tier'] == 'Below fit floor': return '; '.join(parts + ['below the fit floor, not scored'])
    if a['score'] is not None: parts.append(f"score {a['score']:g}")
    parts.append(f"tier {a['tier']}")
    if a['confidence']: parts.append(f"confidence {a['confidence']}")
    if a['tier'] != 'Watch list':
        parts.append('micro-segment ' + after_text(a))
        if a['rec']:
            parts.append(f"action \"{a['action']}\"" + (f", runner-up \"{a['runnerUp']}\"" if a['runnerUp'] else ''))
            parts.append(f"to {a['contact']} by {'nobody (no way to reach them)' if a['reach'] else a['channel']}")
            parts.append('For Review: ' + review_text(a))
    return '; '.join(parts)

# ---------- one workbook per scenario set ----------
FILES = {'TS6': 'TS6_Segment_boundaries.xlsx', 'TS7': 'TS7_Scoring_boundaries.xlsx', 'TS8': 'TS8_Contactability_and_channel.xlsx',
         'TS9': 'TS9_For_Review_routing.xlsx', 'TS10': 'TS10_Rep_ownership.xlsx', 'TS12': 'TS12_Realistic_mixed_list.xlsx'}
NOTES = {
    'TS6': ['Every account here goes through Scoring, which is how micro-segments are normally built: they come from the scored list.',
            'Expect 3 segments: Inquiry 3, Service renewal 4 (Ironbark Refining joins on its next signal), Leadership 3, and 4 exceptions in For Review.',
            'Engagement never forms, although 3 accounts have an Engagement signal: for two of them it is only their second-best signal, and a segment never forms on second-best signals alone.'],
    'TS7': ['Score = fit x timing / 100. Timing = stacked signal weight / 200 x 100. Tier A from 35, B from 20. Fit floor 50.',
            'Bayfront and Driftwood land exactly on the A and B cut-offs; Coldwater and Eastgate are just under. If either cut-off is read as "more than" instead of "at least", those two rows show it.',
            'Rosewood Refining is a known gap: its only inquiry is from a contact who opted out, but it still counts.'],
    'TS8': ['All accounts have a fresh webinar signal so they share the Engagement segment; only the contact details differ.',
            'Hollis and Inglewood are known gaps: see the "Known gap" column. Log them in the Gap Log if the build still behaves this way.'],
    'TS9': ['Three clean items (Alderbank, Beacon Hill, Cobalt Bay) should go ahead on their own; everything else must be in For Review with the reason shown.',
            'The sensitive, claim and brand items are triggered by the wording of the signal detail, which the draft quotes. If the draft wording changes, check those three again.',
            'Marlowe Refining is an inquiry with low confidence: it must be due the same day, not in two days.'],
    'TS10': ['Switch the demo role in the sidebar. The Sales rep role is Rep A (rep.a@client-sample.com, shown as Sofia Ahlgren); the Sales manager sees everyone.',
             'As Sales rep, expect to see only Arden, Brookfield, Carrow and Hartwell. As Sales manager, all 9.',
             'Note: the Users page lists Sofia Ahlgren as sofia.ahlgren@client.com, not rep.a@client-sample.com. Real owner emails will not match the demo rep until the two are joined up.'],
    'TS12': ['This set is for volume and robustness, not one rule per row. Check the counts on this sheet, that every page loads in reasonable time, and that nothing errors.',
             'It is generated with a fixed seed, so rebuilding it gives the same file.'],
}
counts = {}
for st in sc['sets']:
    sid, r = st['id'], scr['sets'][st['id']]
    wb = Workbook()
    sheet(wb, 'Lead template', sc['leadCols'], st['leads'], lead_w, first=True)
    sheet(wb, 'Signal template', sc['sigCols'], st['signals'], sig_w)
    intent = {a['name']: a for a in st['accounts']}
    rows = []
    for a in r['accounts']:
        d = intent.get(a['name'], {})
        rows.append([a['name'], d.get('intent', ''), a['owner'], num(a['fit']), num(a['timing']), num(a['score']), a['tier'], a['signals'], a['confidence'],
                     after_text(a), a['action'], a['runnerUp'], a['contact'], ('None: nobody can be reached' if a['reach'] else a['channel']) if a['rec'] else '',
                     review_text(a), (d.get('expect') or {}).get('gap', '')])
    if sid == 'TS12': rows = [x[:1] + x[2:] for x in rows]
    hdr = ['Account', 'What it tests', 'Owner', 'Fit (0-100)', 'Timing (0-100)', 'Score (0-100)', 'Tier', 'Qualified signals', 'Confidence', 'Micro-segment',
           'Recommended action', 'Runner-up action', 'Recommended contact', 'Channel', 'For Review reasons', 'Known gap (the engine does this today; it should not)']
    wid = [26, 50, 28, 8, 8, 8, 12, 8, 10, 30, 34, 34, 16, 14, 40, 50]
    if sid == 'TS12': hdr, wid = hdr[:1] + hdr[2:], wid[:1] + wid[2:]
    sheet(wb, 'Answer key - accounts', hdr, rows, wid)
    srows = [[s['row'], s['company'], s['type'], s.get('age', ''), s['status'], s.get('code') or '', round(s['weight'], 1) if s.get('weight') is not None else '', s['reason']] for s in r['signals']]
    sheet(wb, 'Answer key - signals', ['Row', 'Account', 'Type of signal', f'Age on {AS_OF} (days)', 'Status', 'Code', 'Weight', 'Reason'], srows, [6, 26, 34, 10, 14, 8, 8, 60])
    stt = r['stats']
    tiers = {}
    for a in r['accounts']: tiers[a['tier']] = tiers.get(a['tier'], 0) + 1
    in_review = sum(1 for a in r['accounts'] if a['rec'] and a['cats'])
    goes = sum(1 for a in r['accounts'] if a['rec'] and not a['cats'])
    counts[sid] = dict(stats=stt, tiers=tiers, segs=r['segments'], exc=r['exceptions'], review=in_review, ahead=goes)
    tier_txt = ', '.join(f'{k} {v}' for k, v in tiers.items())
    seg_txt = ', '.join(f"{s['name']} {s['n']}" for s in r['segments']) or 'none'
    guide(wb, [
        f"{sid} {st['title']}: {len(st['accounts'])} accounts, {len(st['leads'])} contacts, {len(st['signals'])} signals",
        st['purpose'],
        '',
        f"1. Prospecting: upload this file (it reads \"Lead template\"). Expect {stt['rowsRead']} rows read, {stt['accounts']} accounts of which {stt['accountsPass']} pass, "
        f"{stt['contacts']} contacts, {stt['inFinalList']} in the final list. Save the list as \"{sid}\".",
        f'2. Account Scoring: pick the saved list "{sid}", set "Score as of" to {AS_OF}, then upload this same file (it reads "Signal template"). Expect: {tier_txt}.',
        f'3. Micro-segments & NBA: build from this scoring, scoring date {AS_OF}. Expect segments: {seg_txt}; {r["exceptions"]} exception(s).',
        f'4. For Review: expect {in_review} item(s) needing a person, with the reasons in "Answer key - accounts", and {goes} going ahead on their own.',
        '',
        *NOTES[sid],
        '',
        'Compare each account with "Answer key - accounts" and each signal with "Answer key - signals". The answer key was produced by running this file through the POC\'s own rules engine (legacy/demandai-engine.js).',
        f'Dates matter: signal ages are worked out from the scoring date. On any date other than {AS_OF} scores and tiers will differ.',
    ])
    wb.save(os.path.join(outdir, FILES[sid]))

# ---------- TS11: TS1 scored at later dates ----------
wb = Workbook()
srow = []
for x in scr['ts11']:
    t = x['tiers']
    srow.append([x['asOf'], t.get('A', 0), t.get('B', 0), t.get('C', 0), t.get('Watch list', 0), t.get('Below fit floor', 0),
                 ', '.join(f"{s['name']} {s['n']}" for s in x['segments']) or 'none', x['exceptions']])
sheet(wb, 'Answer key - by date', ['Score as of', 'Tier A', 'Tier B', 'Tier C', 'Watch list', 'Below fit floor', 'Segments', 'Exceptions'], srow, [12, 8, 8, 8, 10, 12, 50, 10], first=True)
names = [a['name'] for a in scr['ts11'][0]['accounts']]
by = [{a['name']: a for a in x['accounts']} for x in scr['ts11']]
arow = [[n] + [f"{num(b[n]['score'])} {b[n]['tier']}".strip() + (f" / {b[n]['segment']}" if b[n]['segment'] else '') for b in by] for n in names]
sheet(wb, 'Answer key - accounts', ['Account'] + [x['asOf'] for x in scr['ts11']], arow, [26] + [24] * len(scr['ts11']))
guide(wb, [
    'TS11 Score date: the TS1 files scored on five dates',
    'Use it to check that signals age and fade as they should. There are no new files: use TS1_End_to_end.xlsx.',
    '',
    'In Account Scoring pick the saved list "TS1", set "Score as of" to each date in "Answer key - by date", and upload the TS1 file each time. Then build Micro-segments from that scoring with the same date.',
    'Expect the tier counts and segments in "Answer key - by date", and each account\'s score, tier and segment in "Answer key - accounts".',
    'What to look for: inquiries fade after 30 days and are gone at 90, so the Inquiry segment disappears by 20 Dec; engagement signals go next; end-of-support and project signals last longest.',
    'A scoring date in the future is allowed in the POC, so this can be run any day.',
])
wb.save(os.path.join(outdir, 'TS11_Score_date_answer_key.xlsx'))

# ---------- POC_Test_Cases.xlsx ----------
cases = []
def case(module, area, dataset, title, steps, expected, priority='Medium'):
    cases.append([module, area, dataset, title, steps, expected, priority])
t1 = tsr['ts1']; s1 = t1['stats']
pro = 'Prospecting'
case(pro, 'Upload', 'TS1', 'Upload a clean lead file', 'Upload TS1_End_to_end.xlsx in Prospecting.',
     f"{s1['rowsRead']} rows read, 0 issues, {s1['accounts']} accounts, {s1['accountsPass']} passing, {s1['inFinalList']} contacts in the final list.", 'High')
case(pro, 'Upload', 'TS1', 'Picks the right sheet in a workbook with several', 'Upload TS1_End_to_end.xlsx (it has lead, signal and answer-key sheets).', 'Prospecting reads "Lead template" and ignores the other sheets.', 'High')
case(pro, 'Save list', 'TS1', 'Save the list and find it in Scoring', 'After the upload, save the list as "TS1". Open Account Scoring.', 'The list "TS1" is offered in Account Scoring with 16 accounts.', 'High')
case(pro, 'Download', 'TS1', 'Download the final list', 'Download the final list as CSV and Excel.', 'Both files open, with the 48 contacts and the same columns as on screen.')
case(pro, 'Reload', 'TS1', 'Saved list survives a reload', 'Reload the browser, then open Account Scoring.', 'The list "TS1" is still there (it is kept in the browser).')
t2 = tsr['ts2']; s2 = t2['stats']; cst = {c['row']: c for c in t2['contacts']}
case(pro, 'Data quality', 'TS2', 'Edge-case file totals', 'Upload TS2_Prospecting_edge_cases.xlsx.',
     f"{s2['rowsRead']} rows read, {s2['notLoaded']} not loaded, {s2['merged']} merged, {s2['contacts']} contacts, {s2['accounts']} accounts ({s2['accountsPass']} pass), {s2['inFinalList']} in the final list.", 'High')
for i, (row, tests, intent) in enumerate(ts['ts2']):
    rn = i + 2
    iss = [x for x in t2['issues'] if x['row'] == rn and x['code'] != '—']
    st = cst[rn]['status'] if rn in cst else ('Not loaded' if any(x['severity'] == 'stop' for x in iss) else 'Merged into the earlier row')
    exp = ('Issues: ' + '; '.join(f"{x['field']} {x['code']}" for x in iss) if iss else 'No issue') + f'. Contact: {st}.'
    case(pro, 'Data quality', 'TS2', f'Row {rn}: {intent}', f'In the TS2 upload, find spreadsheet row {rn} in the check pane and the final list.', exp)
sco = 'Account Scoring'
case(sco, 'Scoring', 'TS1', 'Score the end-to-end list', f'Pick list "TS1", set Score as of {AS_OF}, upload TS1_End_to_end.xlsx.',
     'Tier A 4, B 8, C 2, watch list 1 (Penang Refining), below fit floor 1 (Delta Biofuels Trading). Each account as in TS1 "Answer key - accounts".', 'High')
case(sco, 'Scoring', 'TS1', 'Fit breakdown adds up', 'Open Northgate Refining and read the fit breakdown.', 'The six parts (vertical 30, size 20, site 20, relationship 15, account type 10, contacts 5) add up to the fit shown.')
case(sco, 'Scoring', 'TS1', 'Why-now line names the evidence', 'Open any tier A account.', 'The why-now line quotes the strongest signal(s) with the person, and the source and date.')
case(sco, 'Scoring', 'TS1', 'Save and reopen a scoring run', 'Save the scoring run, reload, reopen it.', 'Same tiers and scores as before the reload.')
case(sco, 'Scoring', 'TS1', 'Send to Micro-segments', 'Send the scored list to Micro-segments & NBA.', 'Micro-segments & NBA opens with the TS1 scoring selected.', 'High')
case(sco, 'Setup link', 'TS3', 'A signal type switched off in Setup is rejected', 'In Setup switch off "Content download". Score TS1 with the TS3 file. Switch it back on.', 'Row 14 is rejected as OFF (switched off in Setup).')
for (row, (wst, wcode), intent), s in zip(ts['ts3'], tsr['ts3']['signals']):
    case(sco, 'Signal rules', 'TS3', f"Signal row {s['row']}: {intent}", f'Pick list "TS1", Score as of {AS_OF}, upload TS3_Signal_edge_cases.xlsx; find row {s["row"]}.',
         f"Status {s['status']}" + (f", code {s['code']}" if s.get('code') else '') + (f": {s['reason']}" if s.get('reason') else ''))
SET_MOD = {'TS6': 'Micro-segments & NBA', 'TS7': sco, 'TS8': 'Micro-segments & NBA', 'TS9': 'For Review', 'TS10': 'Roles and visibility'}
SET_AREA = {'TS6': 'Segment rules', 'TS7': 'Fit and tiers', 'TS8': 'Channel and contactability', 'TS9': 'Review reasons', 'TS10': 'Ownership'}
for st in sc['sets']:
    sid = st['id']
    if sid == 'TS12': continue
    got = {a['name']: a for a in scr['sets'][sid]['accounts']}
    c = counts[sid]
    case(SET_MOD[sid], 'Whole set', sid, f"{st['title']}: run the set end to end",
         f'Prospecting: upload {FILES[sid]}, save as "{sid}". Scoring: list "{sid}", Score as of {AS_OF}, same file. Build Micro-segments from that scoring.',
         f"Tiers: {', '.join(f'{k} {v}' for k, v in c['tiers'].items())}. Segments: {', '.join(s['name'] + ' ' + str(s['n']) for s in c['segs']) or 'none'}. "
         f"{c['review']} item(s) in For Review, {c['ahead']} going ahead.", 'High')
    for a in st['accounts']:
        g = got[a['name']]
        gap = a['expect'].get('gap')
        exp = expected_line(g)
        if sid == 'TS10':
            exp = (f"Owner {g['owner'] or '(none)'}. Sales rep (Rep A) " + ('sees it' if a['expect']['repASees'] else 'does not see it') + '; Sales manager sees it.')
        if gap: exp = f'SHOULD: {a["intent"]}. TODAY (known gap): {gap}. Engine result: {expected_line(g)}'
        case(SET_MOD[sid], SET_AREA[sid], sid, f"{a['name']}: {a['intent']}", f'In the {sid} run, open {a["name"]}' + (' as Sales rep, then as Sales manager.' if sid == 'TS10' else '.'),
             exp, 'High' if gap else 'Medium')
nba = 'Micro-segments & NBA'
case(nba, 'Segments', 'TS1', 'Segments from the TS1 scoring', f'Build Micro-segments from the TS1 scoring, scoring date {AS_OF}.', 'Inquiry 3, Modernisation 3, Project 3, Engagement 4; Chao Phraya Refining is an exception.', 'High')
case(nba, 'Segments', 'TS1', 'Fallback shows the runner-up from the strongest signal', 'Open Sumatra Petrochemicals.', 'It sits in Engagement; runner-up is "Offer a renewal review", with the reason that its strongest signal had too few accounts.')
case(nba, 'Draft', 'TS1', 'Draft quotes the signal and has one proof point', 'Open the draft for Northgate Refining.', 'Subject of at most 6 words; body under 120 words; opens on their inquiry; one cited proof point; one question to finish.')
case(nba, 'Draft', 'TS1', 'Wrong contact alternative', 'On Northgate Refining choose another contact.', 'The draft switches to the other person and their own evidence.')
case(nba, 'Upload route (optional)', 'TS4', 'Micro-segments straight from a file', 'Only if lists scored outside the POC are in scope: in Micro-segments & NBA choose "Upload a CSV…", date ' + AS_OF + ', upload TS4.',
     'Inquiry 4, Modernisation 3; 2 rows rejected; Johor Aromatics, Rayong Olefins and Cilacap Refining are exceptions.', 'Low')
rev = 'For Review'
for title, exp, pr in [
        ('Approve an item', 'It leaves the queue as approved and is counted in Today and Analytics.', 'High'),
        ('Reject an item with a reason', 'It leaves the queue; the reason is kept and shows in Analytics.', 'High'),
        ('Edit a draft, then approve', 'The edited text is kept; the checks run again on the edited text.', 'High'),
        ('Send an item back', 'It returns with the "Sent back" reason, due the same day.', 'Medium'),
        ('Assign a micro-segment to an exception', 'The item moves into the chosen segment and the "Segment too small" reason clears.', 'High'),
        ('Add an email for a contact nobody can reach', 'The "Can\'t reach them" reason clears and the draft uses email; the fix shows in Prospecting and Scoring for that contact too.', 'Medium'),
        ('Pause a whole micro-segment (Sales manager)', 'Its items stay in place but cannot be decided until resumed.', 'Medium'),
        ('Spot-check share', 'About 20% of the items that went ahead on their own are drawn for a spot-check, due Friday.', 'Low')]:
    case(rev, 'Actions', 'TS9', title, 'Use the TS9 run in For Review.', exp, pr)
case('Today', 'Counts', 'TS9', 'Today shows what needs a person', 'Open Today after the TS9 run.', f"{counts['TS9']['review']} items waiting, grouped by due date; Marlowe Refining and Fernhill Refining due today.", 'High')
case('Today', 'Counts', 'TS1', 'Today matches For Review after decisions', 'Decide a few TS1 items in For Review, then open Today.', 'Waiting and decided counts match For Review.')
case('Analytics', 'Counts', 'TS1', 'Analytics counts the decisions', 'After deciding items for TS1, open Analytics.', 'Approved, rejected and edited counts match what was done; gates read "too early to call" below 10 decisions.', 'High')
case('Content library', 'Brand guideline', 'TS5', 'Rules read from an uploaded guideline', 'Upload TS5_Brand_guideline.md in Content library > Brand guidelines.', '16 rules found, exactly the list in the TS5 answer key; the "Not rules" lines are ignored.', 'High')
case('Content library', 'Brand guideline', 'TS5', 'A new brand version is used by the checks', 'Publish the TS5 rules as a new brand version; write "synergy" and "Rivalco" into a draft in For Review.', 'Both are flagged (banned word, competitor).')
for x in scr['ts11']:
    t = x['tiers']
    case(sco, 'Signal decay', 'TS11 (TS1 files)', f"Score TS1 as of {x['asOf']}", f"Pick list \"TS1\", Score as of {x['asOf']}, upload TS1_End_to_end.xlsx; build Micro-segments with the same date.",
         f"A {t.get('A', 0)}, B {t.get('B', 0)}, C {t.get('C', 0)}, watch list {t.get('Watch list', 0)}, below fit floor {t.get('Below fit floor', 0)}. "
         f"Segments: {', '.join(s['name'] + ' ' + str(s['n']) for s in x['segments']) or 'none'}.", 'Medium')
c = counts['TS12']; st12 = c['stats']
for mod, title, exp in [
        (pro, 'Volume list loads', f"{st12['rowsRead']} rows read, {st12['accounts']} accounts ({st12['accountsPass']} pass), {st12['contacts']} contacts, {st12['inFinalList']} in the final list; the page stays responsive."),
        (sco, 'Volume list scores', f"Tiers: {', '.join(f'{k} {v}' for k, v in c['tiers'].items())}."),
        (nba, 'Volume list segments', f"Segments: {', '.join(s['name'] + ' ' + str(s['n']) for s in c['segs'])}; {c['exc']} exceptions."),
        (rev, 'Volume list review queue', f"{c['review']} items need a person, {c['ahead']} go ahead on their own."),
        ('Today', 'Volume list in Today', 'Counts match For Review; the page loads without errors.')]:
    case(mod, 'Volume', 'TS12', title, f'Run TS12_Realistic_mixed_list.xlsx through Prospecting, Scoring (as of {AS_OF}) and Micro-segments.', exp, 'Medium')
case(pro, 'Volume', 'Lead_Research_Sample_Dataset.xlsx', 'Large file (7,498 leads)', 'Upload test-data/Lead_Research_Sample_Dataset.xlsx in Prospecting.', 'It loads without the browser freezing; note the time taken.', 'Low')
case('Roles and visibility', 'Role switcher', 'any', 'Each role sees only its pages', 'Switch through each demo role in the sidebar.', 'Each role sees the pages set for it in Users and roles; Sales rep sees only Rep A items; Admin and Sales manager see all items.', 'Medium')

wb = Workbook()
ORDER = [pro, sco, nba, rev, 'Today', 'Analytics', 'Content library', 'Roles and visibility']
cases.sort(key=lambda x: ORDER.index(x[0]))
rows = []
prefix = {pro: 'PR', sco: 'SC', nba: 'NB', rev: 'RV', 'Today': 'TD', 'Analytics': 'AN', 'Content library': 'CL', 'Roles and visibility': 'RL'}
seen = {}
for x in cases:
    seen[x[0]] = seen.get(x[0], 0) + 1
    rows.append([f"{prefix[x[0]]}-{seen[x[0]]:03d}"] + x + ['Not run', '', '', '', ''])
ws = sheet(wb, 'Test cases', ['ID', 'Module', 'Area', 'Dataset', 'Test case', 'Steps', 'Expected result', 'Priority', 'Status', 'Actual result', 'Gap ID', 'Tested on', 'Notes'],
           rows, [8, 18, 18, 12, 50, 50, 60, 9, 11, 40, 9, 11, 30], first=True)
n = len(rows) + 1
dv = DataValidation(type='list', formula1='"Not run,Pass,Fail,Blocked,Not in scope"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'I2:I{n}')
dv2 = DataValidation(type='list', formula1='"High,Medium,Low"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'H2:H{n}')
for col in ('J', 'K', 'L', 'M'):
    for r in range(2, n + 1): ws[f'{col}{r}'].font = Font(name=F, size=10, color='FF0000FF')
for txt, color in (('Pass', 'FFE2F0D9'), ('Fail', 'FFF8D7DA'), ('Blocked', 'FFFFF2CC'), ('Not in scope', 'FFEDEDED')):
    ws.conditional_formatting.add(f'I2:I{n}', FormulaRule(formula=[f'$I2="{txt}"'], fill=PatternFill('solid', fgColor=color)))
ws.auto_filter.ref = f'A1:M{n}'
for r in range(2, n + 1): ws[f'L{r}'].number_format = 'dd-mmm'
# Summary by module, live from the Status column
sm = wb.create_sheet('Summary')
sm.append(['Module', 'Cases', 'Not run', 'Pass', 'Fail', 'Blocked', 'Not in scope', 'Pass rate (of run)'])
for i, m in enumerate(ORDER, 2):
    sm.append([m, f'=COUNTIF(\'Test cases\'!$B:$B,A{i})'] +
              [f'=COUNTIFS(\'Test cases\'!$B:$B,$A{i},\'Test cases\'!$I:$I,"{s}")' for s in ('Not run', 'Pass', 'Fail', 'Blocked', 'Not in scope')] +
              [f'=IFERROR(D{i}/(D{i}+E{i}+F{i}),"")'])
t = len(ORDER) + 2
sm.append(['Total'] + [f'=SUM({c}2:{c}{t - 1})' for c in 'BCDEFG'] + [f'=IFERROR(D{t}/(D{t}+E{t}+F{t}),"")'])
for row in sm.iter_rows():
    for cl in row:
        cl.font = Font(name=F, size=10, bold=cl.row in (1, t), color='FFFFFFFF' if cl.row == 1 else 'FF000000'); cl.border = B
        if cl.row == 1: cl.fill = HDR
        if cl.column == 8 and cl.row > 1: cl.number_format = '0%'
sm.column_dimensions['A'].width = 22
for c in 'BCDEFGH': sm.column_dimensions[c].width = 12
ds = [['TS1', 'TS1_End_to_end.xlsx', 'The list carried through every module'], ['TS2', 'TS2_Prospecting_edge_cases.xlsx', 'Prospecting data rules, one per row'],
      ['TS3', 'TS3_Signal_edge_cases.xlsx', 'Signal rules, one per row (uses the TS1 list)'], ['TS4', 'TS4_Microsegments_from_file.xlsx', 'Optional: the NBA "Upload a CSV" route for lists scored outside the POC'],
      ['TS5', 'TS5_Brand_guideline.md', 'Brand rules read from a guideline'], ['TS6', FILES['TS6'], 'Which micro-segments form, fallbacks and exceptions'],
      ['TS7', FILES['TS7'], 'Fit floor, rubric parts, tier cut-offs, decay, stacking, knock-out'], ['TS8', FILES['TS8'], 'Channel choice and who cannot be reached'],
      ['TS9', FILES['TS9'], 'Each For Review reason and due date'], ['TS10', FILES['TS10'], 'Which role sees which item'],
      ['TS11', 'TS11_Score_date_answer_key.xlsx', 'TS1 scored on five dates: signals fading'], ['TS12', FILES['TS12'], 'A realistic 60-account list for volume and robustness']]
sheet(wb, 'Datasets', ['Set', 'File', 'What it is for'], ds, [8, 40, 70])
guide(wb, [
    f'POC test cases: {len(rows)} cases across {len(ORDER)} modules',
    'Each case names the dataset to use, the steps and the expected result. Expected results were produced by running the dataset through the POC\'s own rules engine.',
    '',
    'How to use it in the test window:',
    '1. Filter "Test cases" by Module for the day\'s module (see the Scrum Tracker). Run the High-priority cases first.',
    '2. Set Status to Pass, Fail, Blocked or Not in scope, and write what actually happened in "Actual result" when it differs.',
    '3. For every Fail, add a row in the Gap Log of the project plan and put its ID in "Gap ID".',
    '4. "Summary" counts cases by module and status, for the Scrum Tracker columns "Cases planned / run / passed".',
    '',
    f'Score as of: always {AS_OF} unless the case gives another date (TS11).',
    'Cases whose expected result starts with "SHOULD" are known gaps: the engine does something else today. They are expected to fail until fixed; log them once.',
    'Micro-segments normally come from Scoring, so TS6 to TS10 all go Prospecting > Scoring > Micro-segments. TS4 is only for the separate "Upload a CSV" route; mark it Not in scope if that route is not part of the POC.',
])
wb.move_sheet('How to use', offset=-(len(wb.sheetnames) - 1))
wb.save(os.path.join(outdir, 'POC_Test_Cases.xlsx'))
print(len(rows), 'test cases;', ', '.join(f'{m} {seen[m]}' for m in ORDER))
