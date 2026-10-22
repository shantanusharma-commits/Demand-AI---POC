"""Writes the POC test-set workbooks from make_inputs.py's rows and run_engine.js's results.
Usage: python3 build_poc_test_sets.py <inputs.json> <results.json> <output folder>"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

inp, res, outdir = json.load(open(sys.argv[1])), json.load(open(sys.argv[2])), sys.argv[3]
F = 'Arial'
thin = Side(style='thin', color='FFBFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F1F1F')
AS_OF = inp['asOf']

def sheet(wb, title, header, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=F, size=10, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'A2'
    return ws

def guide(wb, lines):
    ws = wb.create_sheet('How to use')
    ws.column_dimensions['A'].width = 120
    for i, l in enumerate(lines, 1):
        c = ws.cell(i, 1, l); c.font = Font(name=F, size=11 if i == 1 else 10, bold=(i == 1)); c.alignment = Alignment(wrap_text=True, vertical='top')

lead_w = [12, 12, 30, 10, 34, 8, 40, 26, 26, 12, 14, 18, 14, 10, 15, 9, 9, 20, 18, 10, 8, 18, 8, 30]
sig_w = [26, 36, 34, 12, 50, 16, 14]

# ---------- TS1 ----------
wb = Workbook()
sheet(wb, 'Lead template', inp['leadCols'], inp['ts1'], lead_w, first=True)
sheet(wb, 'Signal template', inp['sigCols'], inp['ts1sig'], sig_w)
t1 = res['ts1']
rows = []
for a in t1['accounts']:
    if a['excluded']: after = 'Excluded: ' + a['excludedReason']
    elif a['tier'] == 'Below fit floor': after = 'Below fit floor (fit under 50): not scored, not passed to Micro-segments'
    elif a['tier'] == 'Watch list': after = 'No signal: watch list, not passed to Micro-segments'
    elif a['segment']: after = f"Segment: {a['segment']}"
    else: after = 'No segment: goes to For Review as an exception'
    rows.append([a['name'], f"{a['vertical']} ({a['level']})", 'In the list, all 3 contacts in the final list', a['fit'], a.get('timing', ''), a.get('score', ''), a['tier'],
                 a['signals'], after, a['action'] or '', a['runnerUp'] or '', a['exception'] or ''])
sheet(wb, 'Answer key - accounts', ['Account', 'Vertical', 'Prospecting result', 'Fit (0-100)', 'Timing (0-100)', 'Score (0-100)', 'Tier', 'Qualified signals',
      'Micro-segments result', 'Recommended action', 'Runner-up action', 'Exception reason'], rows, [26, 18, 26, 9, 9, 9, 14, 9, 34, 34, 34, 34])
srows = []
for s, src in zip(t1['signals'], inp['ts1sig']):
    srows.append([s['row'], src[0], src[2], src[3], s['age'], s['status'], s['code'], round(s['weight'], 1) if s['weight'] is not None else ''])
sheet(wb, 'Answer key - signals', ['Row', 'Account', 'Type of signal', 'Event date', 'Age on ' + AS_OF + ' (days)', 'Status', 'Code', 'Weight'], srows, [6, 26, 34, 12, 12, 14, 8, 8])
guide(wb, [
    'TS1 End-to-end set: 16 accounts, 48 contacts, 17 signals, all clean',
    'Use it to walk one list through every module: Prospecting > Account Scoring > Micro-segments & NBA > For Review > Today > Analytics.',
    '',
    '1. Prospecting: upload this file. It reads the "Lead template" sheet. Expect 48 rows read, 0 issues, 16 accounts, all 16 passing, 48 contacts in the final list. Save the list as "TS1".',
    f'2. Account Scoring: pick the saved list "TS1", set "Score as of" to {AS_OF}, then upload this same file (it reads the "Signal template" sheet).',
    '   Expect: 4 accounts in tier A, 8 in B, 2 in C, 1 on the watch list (Penang Refining), 1 below the fit floor (Delta Biofuels Trading). Compare each account with "Answer key - accounts".',
    f'3. Micro-segments & NBA: build from the TS1 scoring, scoring date {AS_OF}. Expect 4 segments: Inquiry 3, Modernisation 3, Project 3, Engagement 4.',
    '   Sumatra Petrochemicals joins Engagement on its second signal (only it has a service renewal), so its runner-up is the renewal action.',
    '   Chao Phraya Refining has only a leadership change, which no other account shares: no segment, so it goes to For Review as an exception.',
    '4. For Review: the Chao Phraya exception must be there. Approve, reject and hold a few items; then check Today and Analytics count the same decisions.',
    '',
    'Dates matter: signal ages are worked out from the scoring date. On any other date the timing, scores and tiers will differ from the answer key.',
    'The answer key was produced by running this file through the POC\'s own rules engine (legacy/demandai-engine.js) and checked against what each row was designed to show.',
])
wb.save(os.path.join(outdir, 'TS1_End_to_end.xlsx'))

# ---------- TS2 ----------
wb = Workbook()
sheet(wb, 'Lead template', inp['leadCols'], [x[0] for x in inp['ts2']], lead_w, first=True)
t2 = res['ts2']; cst = {c['row']: c for c in t2['contacts']}
rows = []
for i, (row, tests, intent) in enumerate(inp['ts2']):
    rn = i + 2
    iss = [x for x in t2['issues'] if x['row'] == rn]
    flags = '; '.join(f"{x['field']}: {x['code']} {x['issue']}" for x in iss if x['code'] not in ('—',)) or 'None'
    infos = '; '.join(f"{x['field']}: {x['issue']}" for x in iss if x['code'] == '—') or ''
    st = cst[rn]['status'] if rn in cst else ('Not loaded' if any(x['severity'] == 'stop' for x in iss) else 'Merged into the earlier row')
    rows.append([rn, row[7] or '(blank)', ' '.join(v for v in row[:2] if v) or '(blank)', intent, flags, infos, st])
sheet(wb, 'Answer key - rows', ['Row', 'Account', 'Contact', 'What the row tests', 'Expected issues (field: code)', 'Expected notes (no code)', 'Expected contact status'],
      rows, [6, 26, 18, 44, 54, 34, 40])
arows = [[a['name'], a['vertical'], a['confidence'], 'Yes' if a['needsReview'] else 'No', a['excludedReason'] or 'Passes'] for a in t2['accounts']]
sheet(wb, 'Answer key - accounts', ['Account', 'Vertical', 'Classification confidence', 'Held for review', 'Knock-out result'], arows, [28, 16, 14, 12, 44])
st2 = t2['stats']
guide(wb, [
    'TS2 Prospecting edge cases: 32 rows, one data problem per row',
    'Use it to check each data-quality rule in Prospecting on its own, so a failure points at one rule.',
    '',
    f"Upload in Prospecting. Expect: {st2['rowsRead']} rows read, {st2['notLoaded']} not loaded, {st2['merged']} merged as duplicates, {st2['contacts']} contacts, "
    f"{st2['accounts']} accounts of which {st2['accountsPass']} pass, {st2['inFinalList']} contacts in the final list.",
    'Then open the check pane and the final list and compare each row with "Answer key - rows" (row numbers are spreadsheet rows), and each account with "Answer key - accounts".',
    '',
    'Note on row 33 (country written "SG"): the POC matches country names exactly, so "SG" fails the region check and the account is excluded.',
    'That is what the rules do today. If the business expects "SG" to be read as Singapore, log it as a gap.',
    'Do not score this list: it is for the Prospecting checks only. TS1 is the set to carry forward.',
])
wb.save(os.path.join(outdir, 'TS2_Prospecting_edge_cases.xlsx'))

# ---------- TS3 ----------
wb = Workbook()
sheet(wb, 'Signal template', inp['sigCols'], [x[0] for x in inp['ts3']], sig_w, first=True)
rows = []
for (row, (wst, wcode), intent), s in zip(inp['ts3'], res['ts3']['signals']):
    rows.append([s['row'], row[0] or '(blank)', row[2], row[3] or '(blank)', intent, s['status'], s['code'], s['reason'] or s['assigned']])
sheet(wb, 'Answer key - signals', ['Row', 'Account', 'Type of signal', 'Event date', 'What the row tests', 'Expected status', 'Expected code', 'Expected reason'],
      rows, [6, 26, 30, 12, 44, 18, 9, 60])
guide(wb, [
    'TS3 Signal edge cases: 16 signal rows, one problem per row',
    'Use it to check each signal rule in Account Scoring on its own.',
    '',
    f'1. Run TS1 through Prospecting first and save the list as "TS1". 2. In Account Scoring pick the "TS1" list, set "Score as of" to {AS_OF}, and upload this file.',
    '3. Open the unused-rows drawer and the signal evidence, and compare each row with "Answer key - signals".',
    'Also expect: Penang Refining excluded by the knock-out (pilot product installed), and new.person@kestrelpetchem.com listed as a suggested new contact.',
    'Optional Setup test: switch off "Content download" in Setup and score again. Row 14 should then be rejected as OFF (switched off in Setup). Switch it back on afterwards.',
])
wb.save(os.path.join(outdir, 'TS3_Signal_edge_cases.xlsx'))

# ---------- TS4 ----------
wb = Workbook()
sheet(wb, 'Prospects with signals', inp['segCols'], [x[0] for x in inp['ts4']], [10, 12, 24, 34, 8, 36, 24, 16, 12, 8, 18, 7, 34, 12, 40, 8, 30, 14, 9, 14], first=True)
t4 = res['ts4']; amap = {a['name']: a for a in t4['accounts']}
rows = []
for i, (row, intent) in enumerate(inp['ts4']):
    rn = i + 2; a = amap.get(row[6])
    iss = '; '.join(f"{x['field']}: {x['code']} {x['issue']}" for x in t4['issues'] if x['row'] == rn) or 'None'
    if a:
        result = f"Segment: {a['segment']}" if a['segment'] else 'No segment: ' + a['exception']
        rows.append([rn, row[6], intent, iss, a['score'], 'From the file' if a['scoreFrom'] == 'file' else 'Worked out from the signal', result])
    else:
        rows.append([rn, row[6], intent, iss, '', '', 'Row not used'])
sheet(wb, 'Answer key', ['Row', 'Account', 'What the row tests', 'Expected issues', 'Expected score', 'Score source', 'Expected result'], rows, [6, 24, 50, 50, 9, 22, 56])
guide(wb, [
    'TS4 Micro-segments from a file: 12 prospects already scored outside the POC',
    'Use it to test the second way into Micro-segments & NBA: uploading a list with its signals, without going through Prospecting and Scoring.',
    'Optional: micro-segments normally come from Scoring (TS1, TS6 to TS10 test that). Run this set only if lists scored outside the POC are in scope.',
    '',
    f'In Micro-segments & NBA choose to upload a file, set the scoring date to {AS_OF}, and upload this file.',
    f"Expect {t4['stats']['rowsRead']} rows read, {t4['stats']['rejected']} rejected, {t4['stats']['accounts']} accounts, and 2 segments: Inquiry 4 and Modernisation 3.",
    'Johor Aromatics, Rayong Olefins and Cilacap Refining get no segment, because fewer than 3 accounts share their strongest signal. They go to For Review.',
    'Compare each row with "Answer key".',
])
wb.save(os.path.join(outdir, 'TS4_Microsegments_from_file.xlsx'))

# ---------- TS5 ----------
open(os.path.join(outdir, 'TS5_Brand_guideline.md'), 'w').write(inp['ts5'])
wb = Workbook()
sheet(wb, 'Answer key', ['Kind of rule', 'Value'], [list(x) for x in inp['ts5expect']], [22, 40], first=True)
guide(wb, [
    'TS5 Brand guideline: a Markdown file for the Content library',
    'In Content library > Brand guidelines, upload TS5_Brand_guideline.md. Expect 16 rules under "Rules found in the document", exactly the list in "Answer key".',
    'The two sentences under "Not rules" must not be picked up.',
    'Then add a few rules to the draft, save it as a new brand version, and check a message in For Review or Content studio flags a banned word (e.g. "synergy") and a competitor (e.g. "Rivalco").',
])
wb.save(os.path.join(outdir, 'TS5_Brand_guideline_answer_key.xlsx'))
print('written to', outdir)
