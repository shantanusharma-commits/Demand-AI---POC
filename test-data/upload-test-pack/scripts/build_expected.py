"""Writes Expected_Results.xlsx for the upload test pack from make_pack.py's manifest and run_pack.js's results.
Usage: python3 build_expected.py <manifest.json> <results.json> <pack folder>"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

man, res, pack = json.load(open(sys.argv[1])), json.load(open(sys.argv[2])), sys.argv[3]
F = 'Arial'
thin = Side(style='thin', color='FFBFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='FF1F1F1F')
BAND = PatternFill('solid', fgColor='FFF2F2F2')
STATUS = '"Not run,Pass,Fail,Blocked,Not in scope"'

def sheet(wb, title, header, rows, widths, first=False, status_col=None, band_col=None):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(name=F, size=10, bold=True, color='FFFFFFFF'); c.fill = HDR; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center')
    prev, shade = None, False
    for r in rows:
        ws.append(r)
        if band_col is not None:
            if r[band_col] != prev: shade, prev = not shade, r[band_col]
            if shade:
                for c in ws[ws.max_row]: c.fill = BAND
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'B2'
    n = ws.max_row
    if status_col and n > 1:
        dv = DataValidation(type='list', formula1=STATUS, allow_blank=True); ws.add_data_validation(dv); dv.add(f'{status_col}2:{status_col}{n}')
        for r in range(2, n + 1):
            ws[f'{status_col}{r}'] = ws[f'{status_col}{r}'].value or 'Not run'
            for col in range(ws[f'{status_col}{r}'].column, len(header) + 1): ws.cell(r, col).font = Font(name=F, size=10, color='FF0000FF')
        for txt, color in (('Pass', 'FFE2F0D9'), ('Fail', 'FFF8D7DA'), ('Blocked', 'FFFFF2CC'), ('Not in scope', 'FFEDEDED')):
            ws.conditional_formatting.add(f'{status_col}2:{status_col}{n}', FormulaRule(formula=[f'${status_col}2="{txt}"'], fill=PatternFill('solid', fgColor=color)))
    ws.auto_filter.ref = ws.dimensions
    return ws

def counts(f):
    r = res[f['id']]
    if 'stats' in r and f['kind'] != 'pair':
        s = r['stats']; i = s['issuesBySeverity']
        if i['stop'] and not s['rowsRead']: return 'File not processed (stop): ' + next(x['issue'] for x in r['issues'] if x['severity'] == 'stop')
        return (f"{s['rowsRead']} rows read; {s['notLoaded']} not loaded; {s['merged']} merged; {s['accounts']} accounts, {s['accountsPass']} pass; "
                f"{s['inFinalList']} in the final list. Issues: {i['stop']} stop, {i['flag']} flagged, {i['info']} notes.")
    if f['kind'] == 'pair':
        return (f"Tiers: {', '.join(f'{k} {v}' for k, v in r['tiers'].items())}. Segments: {', '.join(x['name'] + ' ' + str(x['n']) for x in r['segments'])}; "
                f"{r['exceptions']} exception(s).")
    out = []
    for run in r['runs']:
        stop = [x for x in run['fileIssues'] if x['severity'] == 'stop']
        if stop: out.append(f"{run['asOf']}: file not processed: {stop[0]['issue']}"); continue
        st = {}
        for s in run['signals']:
            k = s['status'] + (f" {s['code']}" if s['status'] == 'Rejected' else '')
            st[k] = st.get(k, 0) + 1
        out.append(f"{run['asOf']}: {run['rowsRead']} rows: " + ', '.join(f'{v} {k}' for k, v in st.items()) +
                   f". Tiers: {', '.join(f'{k} {v}' for k, v in run['tiers'].items())}.")
    return '\n'.join(out)

# ---------- Files ----------
rows = []
for n, f in enumerate(man['files'], 1):
    rows.append([f['id'], f['folder'], f['file'], f['title'], f['why'], f['steps'], ', '.join(f.get('asOf', [])) or 'n/a', f['expect'], counts(f), '', '', ''])
wb = Workbook()
sheet(wb, 'Files', ['ID', 'Folder', 'File', 'Use case', 'Why it matters', 'How to run it', 'Score as of', 'What you should see', 'Exact counts (from the POC rules)',
      'Status', 'What actually happened', 'Gap ID'], rows, [6, 18, 34, 30, 40, 44, 12, 50, 50, 11, 34, 9], first=True, status_col='J')

# ---------- Prospecting rows and accounts ----------
prow, pacc = [], []
for f in man['files']:
    if not (f['kind'].startswith('lead') or f['kind'] == 'base'): continue
    r = res[f['id']]; cmap = {c['row']: c for c in r['contacts']}
    for x in f['rows']:
        iss = [i for i in r['issues'] if i['row'] == x['row']]
        flags = '\n'.join(f"{i['field']}: {i['code']} {i['issue']}" for i in iss if i['code'] != '—') or 'None'
        notes = '\n'.join(f"{i['field']}: {i['issue']}" for i in iss if i['code'] == '—')
        c = cmap.get(x['row'])
        status = c['status'] if c else ('Not loaded' if any(i['severity'] == 'stop' for i in iss) else
                                         'Merged into an earlier row' if any(i['code'] == 'D5' for i in iss) else 'Not read')
        prow.append([f['id'], x['row'], x['tests'], flags, notes, status, c['persona'] if c else '', c['reachableBy'] if c else '', 'Not run', ''])
    for a in r['accounts']:
        pacc.append([f['id'], a['name'], 'Excluded: ' + a['excludedReason'] if a['excluded'] else 'Passes', f"{a['vertical']} ({a['level']})", a['confidence'],
                     'Yes' if a['needsReview'] else 'No', a['owner'] or '(none)'])
sheet(wb, 'Prospecting rows', ['File', 'Row', 'What the row tests', 'Expected flags (field: code issue)', 'Expected notes', 'Expected contact status', 'Persona',
      'Reachable by', 'Status', 'What actually happened'], prow, [6, 6, 46, 46, 34, 40, 10, 16, 11, 30], status_col='I', band_col=0)
sheet(wb, 'Prospecting accounts', ['File', 'Account', 'Region and suppression checks', 'Vertical', 'Classification confidence', 'Held for review', 'Owner'],
      pacc, [6, 28, 44, 24, 12, 10, 30], band_col=0)

# ---------- Signal rows and Scoring accounts ----------
srow, sacc = [], []
for f in man['files']:
    if not f['kind'].startswith('signal'): continue
    tests = {x['row']: x['tests'] for x in f['rows']}
    for run in res[f['id']]['runs']:
        for s in run['signals']:
            why = s['reason'] or s['assigned']
            srow.append([f['id'], run['asOf'], s['row'], s['company'] or '(blank)', s['type'] or '(blank)', tests.get(s['row'], ''), s['status'], s['code'] or '',
                         why, s['person'] or '', '' if s['age'] is None else s['age'], '' if s['weight'] is None else s['weight'], 'Not run', ''])
        touched = {s['company'].lower().split('  ')[0].strip() for s in run['signals']}
        for a in run['accounts']:
            if f['folder'].startswith('2-') or a['signals'] or a['name'].lower() in touched:
                sacc.append([f['id'], run['asOf'], a['name'], '' if a['fit'] is None else a['fit'], '' if a['timing'] is None else a['timing'],
                             '' if a['score'] is None else a['score'], a['tier'], a['signals'], a['confidence'] or a['reason']])
sheet(wb, 'Signal rows', ['File', 'Score as of', 'Row', 'Company', 'Type of signal', 'What the row tests', 'Expected status', 'Code', 'Reason, or who it was given to',
      'Person', 'Age (days)', 'Weight', 'Status', 'What actually happened'], srow, [6, 11, 5, 22, 30, 46, 16, 6, 50, 14, 8, 8, 11, 30], status_col='M', band_col=0)
sheet(wb, 'Scoring accounts', ['File', 'Score as of', 'Account', 'Fit (0-100)', 'Timing (0-100)', 'Score (0-100)', 'Tier', 'Qualified signals', 'Confidence, or why not scored'],
      sacc, [6, 11, 28, 9, 9, 9, 14, 9, 44], band_col=0)

# ---------- Summary ----------
sm = wb.create_sheet('Summary')
sm.append(['Folder', 'Files', 'Not run', 'Pass', 'Fail', 'Blocked', 'Not in scope'])
folders = sorted({f['folder'] for f in man['files']})
for i, fo in enumerate(folders, 2):
    sm.append([fo, f'=COUNTIF(Files!$B:$B,A{i})'] + [f'=COUNTIFS(Files!$B:$B,$A{i},Files!$J:$J,"{s}")' for s in ('Not run', 'Pass', 'Fail', 'Blocked', 'Not in scope')])
t = len(folders) + 2
sm.append(['Total'] + [f'=SUM({c}2:{c}{t - 1})' for c in 'BCDEFG'])
for row in sm.iter_rows():
    for c in row:
        c.font = Font(name=F, size=10, bold=c.row in (1, t), color='FFFFFFFF' if c.row == 1 else 'FF000000'); c.border = B
        if c.row == 1: c.fill = HDR
sm.column_dimensions['A'].width = 30

# ---------- How to use ----------
ws = wb.create_sheet('How to use', 0)
ws.column_dimensions['A'].width = 125
lines = [
    'Upload test pack: drop each file into the POC as it is',
    f"{len(man['files'])} files, one use case each. Nothing to edit: drag a file onto the upload box and compare what the POC shows with this workbook.",
    '',
    '1-Prospecting (P01-P22): drop each file on the Prospecting upload box. Compare the check pane and the final list with "Prospecting rows" and "Prospecting accounts".',
    '   To carry a list on into Scoring, use its own signal file, named the same with _SIGNALS on the end (P01_clean_list.xlsx > P01_clean_list_SIGNALS.xlsx).',
    '   In Account Scoring: "Upload a prospect file" = the P list, then drop its _SIGNALS file on box 2. P16, P19 and P20 load nothing, so they have no signal file.',
    '2-Account-Scoring (S01-S07) and 3-Signal-rules (G01-G14): these all use one prospect list, S00.',
    '   Once: Account Scoring > New scoring > 1 · Prospect list > "Upload a prospect file" > drop S00_prospect_list_for_scoring_and_signals.csv. It is saved as a list.',
    f"   Then for each S or G file: choose that S00 list, set \"Score as of\" to {man['asOf']}, and drop the file on box 2. Compare with \"Signal rows\" and \"Scoring accounts\".",
    '   S01 is also scored on two later dates (shown in the "Score as of" column) to see signals fade.',
    '4-Micro-segments-and-Review (N01-N04): each file holds a "Lead template" and a "Signal template" sheet. Upload it as the prospect list, then drop the same file as the signals; then send the scoring to Micro-segments & NBA and open For Review.',
    '   Their per-account answer keys are in ../poc-test-sets (TS6, TS8, TS9, TS10).',
    '',
    f"Always set Score as of to {man['asOf']} unless the row gives another date: signal ages are counted from that date.",
    'Mark each file Pass, Fail, Blocked or Not in scope in "Files" (and row by row if you want). For every Fail, add a row to the Gap Log and put its ID in "Gap ID".',
    'Rows marked "known gap" or "SHOULD" describe what the business would expect; the POC does something else today. They are expected to fail until fixed.',
    '',
    'Every expected result was produced by reading the file exactly as the POC does and running it through the POC\'s own rules (legacy/demandai-engine.js).',
    'All companies, people and domains are invented.',
]
for i, l in enumerate(lines, 1):
    c = ws.cell(i, 1, l); c.font = Font(name=F, size=11 if i == 1 else 10, bold=(i == 1)); c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(os.path.join(pack, 'Expected_Results.xlsx'))
print('rows: files', len(rows), 'prospecting', len(prow), 'signals', len(srow), 'scoring accounts', len(sacc))
